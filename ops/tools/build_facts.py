"""Сборка фактов X5 (data/facts/*.json, data/calendar.json) из первички.

Запуск из корня репозитория (нужен openpyxl):

    python -B ops/tools/build_facts.py                 # переписать data/facts и data/calendar.json
    python -B ops/tools/build_facts.py --out DIR       # записать в DIR/data/... (сверка)

Читает databook X5 (два файла), МСФО, пресс-релизы и трейдинг-апдейты из папки
первички (`--primary` или X5_PRIMARY; по умолчанию ../x5-850-handoff/reference/primary
рядом с репозиторием), реестр облигаций из research/facts папки передачи, аналогов — из
databook Ленты/Fix Price/О'Кей и фактов модели Магнита 850oa (только чтение; соседние
папки рабочего каталога, `--workspace` или X5_WORKSPACE). Каждое число — узел
{"v", "src"} или {"v", "calc"}; нераскрытое — {"v": null, "calc": "почему нет числа"}. Деньги — млрд ₽, площадь — тыс. м²,
акции — млн шт. Файлы пишутся в UTF-8 с переводом строки LF (как их хранит git).
actuals.json (факты журнала) ведёт человек: сборщик пишет его шаблон, только если файла нет,
и не перезаписывает. Проверка «пересборка = data/facts байт в байт» (кроме actuals.json) —
tests/test_facts.py::test_builder_reproduces_facts (метка primary: идёт при первичке рядом).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

# openpyxl импортируется там, где читаются книги Excel: sources(), missing() и write() работают
# и без него (tests/conftest.py решает по ним, пропустить ли тест с меткой primary).

REPO = Path(__file__).resolve().parents[2]

DB26 = "financial_and_operating_results_q2_2026.xlsx"
DB24 = "financial_and_operating_results_q1_2024.xlsx"
FP = "peer-fixprice-pjsc-databook-6m2026.xlsx"

IFRS26 = "finansovaya-otchetnost_pao-kcz-iks5_6m2026.pdf"
IFRS25H1 = "finansovaya-otchetnost_pao-kcz-iks5_6m2025.pdf"
IFRS25 = "auditorskoe-zaklyuchenie-i-finansovaya-otchetnost_pao-kcz-iks-5_2025.pdf"
IFRS24 = "ifrs-fy2024_godovaya-konsolidirovannaya-finansovaya-otchetnost.pdf"
PR226 = "x5_q2_2026_financial_results_rus.pdf"
PR425 = "x5_q4_2025_financial_results_rus.pdf"
ER6M26 = "oeczb-pao-kcz-za-6m2026_disclosure.pdf"
AR25 = "x5-ar25.pdf"

AS_OF = "2026-06-30"


# ------------------------------------------------------------------ узлы

def V(v, src=None, calc=None, **extra):
    """Узел факта: значение + источник и/или формула. Нераскрытое — null с пояснением
    в calc (data/facts/SCHEMA.md); без пояснения — «не раскрыто»."""
    if v is None and not (src or calc):
        calc = "не раскрыто"
    node = {"v": v}
    if src:
        node["src"] = src
    if calc:
        node["calc"] = calc
    node.update(extra)
    # То же правило, что у загрузки (model/facts.py::has_source): источник — текст хотя бы
    # с одной буквой или цифрой; пробел, «?», «—» источником не считаются.
    if v is not None and not any(isinstance(x, str) and any(ch.isalnum() for ch in x)
                                 for x in (src, calc)):
        raise ValueError(f"узел без src и calc (или они пусты по смыслу): {v}")
    return node


def bn(x_mn):
    """млн ₽ → млрд ₽ (3 знака = точность млн)."""
    return None if x_mn is None else round(x_mn / 1000.0, 6)


def ru(x, n=3):
    """Число с запятой для текстов calc."""
    return f"{x:.{n}f}".replace(".", ",")


def r6(x):
    return None if x is None else round(x, 6)


# ------------------------------------------------------------------ книги Excel

class Book:
    """Книга Excel с картой столбцов по заголовкам строки 5 (первое вхождение = IAS 17)."""

    def __init__(self, path: Path, name: str):
        import openpyxl

        self.name = name
        self.wb = openpyxl.load_workbook(path, data_only=True)

    def ws(self, sheet):
        return self.wb[sheet]

    def val(self, sheet, addr):
        v = self.wb[sheet][addr].value
        if isinstance(v, str):
            s = v.replace("\xa0", "").replace(" ", "").replace(",", ".")
            if s.endswith("%"):
                return float(s[:-1]) / 100.0
            try:
                return float(s)
            except ValueError:
                return None
        return v

    def colmap(self, sheet, row=5, prefer_adj_row=None):
        """{заголовок: буква столбца}; при повторе берётся первый (IAS 17), либо «Скорр.»."""
        from openpyxl.utils import get_column_letter

        ws = self.wb[sheet]
        out = {}
        for cell in ws[row]:
            key = cell.value
            if key is None:
                continue
            if hasattr(key, "strftime"):
                key = key.strftime("%Y-%m-%d")
            key = str(key).strip()
            col = get_column_letter(cell.column)
            if key not in out:
                out[key] = col
            elif prefer_adj_row and str(ws[f"{col}{prefer_adj_row}"].value or "").startswith("Скорр"):
                out[key] = col
        return out

    def ref(self, sheet, *addrs):
        return f"{self.name} › {sheet}!" + " + ".join(addrs)


def q_new(q, y):
    return f"{q} КВ. {y}"


def q_old(q, y):
    return f"Q{q} {y}"


# ------------------------------------------------------------------ источники

def sha256_file(path: Path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sources(workspace=None, primary=None) -> dict:
    """Пути к входам сборки: рабочий каталог (папка с репозиторием и соседями) и первичка."""
    ws = Path(workspace or os.environ.get("X5_WORKSPACE") or REPO.parent)
    handoff = ws / "x5-850-handoff"
    return {
        "primary": Path(primary or os.environ.get("X5_PRIMARY") or handoff / "reference" / "primary"),
        "debt_research": handoff / "research" / "facts" / "debt_register.json",
        "magnit": ws / "magnit-850oa",
        "lenta_db": ws / "lenta-850-handoff" / "reference" / "primary" / "Lenta_Q22026_DATABOOK.xlsx",
    }


def missing(src: dict) -> list:
    """Входы, которых нет на диске (сборка без них невозможна)."""
    need = [src["primary"] / DB26, src["primary"] / DB24, src["primary"] / FP, src["debt_research"],
            src["magnit"] / "data" / "facts" / "history.json", src["lenta_db"]]
    return [str(p) for p in need if not p.exists()]


# ------------------------------------------------------------------ история сети до 31.12.2023

NETWORK_SHEET = "data/assumptions/evidence/network-revenue"
TU_2022_2023 = REPO / NETWORK_SHEET / "inputs" / "tu_2022_2023.json"


def network_history_est(b24, orq: dict, area_end: dict, area_by_fmt: dict):
    """Оценки листа «Сеть и выручка» для индекса эффективной площади (docs/MODEL.md §4.1):
    площадь на конец 2022H2 и 2023H1, валовые открытия и закрытия площади 2022H1–2023H2.

    Правило — `data/assumptions/evidence/network-revenue/common.py::network_history`, числа те же
    (сверка — до 1e-6 тыс. м²): databook «как отчитано» (`Operating Results_Q`, три формата) плюс
    площадь «Красного Яра»/«Слаты» (в сопоставимом базисе они внутри «Пятёрочки»); магазины,
    открытые и закрытые по кварталам, — ручная выписка трейдинг-апдейтов 4 кв. 2022 – 4 кв. 2023
    листа (`inputs/tu_2022_2023.json`, со строками текстов первички). Возвращает три словаря
    узлов: площадь, открытия, закрытия (полугодие → узел)."""
    OQ = "Operating Results_Q"
    space_row = {"pyaterochka": 70, "perekrestok": 71, "chizhik": 73}
    stores_row = {"pyaterochka": 62, "perekrestok": 63, "chizhik": 65}
    name = {"pyaterochka": "П", "perekrestok": "Пер", "chizhik": "Ч"}
    fmts3 = ("pyaterochka", "perekrestok", "chizhik")
    sheet = f"оценка листа {NETWORK_SHEET}, common.py::network_history"
    tu = json.loads(TU_2022_2023.read_text(encoding="utf-8"))
    T = tu["quarters"]

    def cell(q, row):
        col = orq[q_old(int(q[-1]), int(q[:4]))]
        return b24.val(OQ, f"{col}{row}"), f"{col}{row}"

    def space(q, f):
        return cell(q, space_row[f])[0]

    def refs(q, rows):
        return b24.ref(OQ, *(cell(q, r)[1] for r in rows))

    def avg_store(q, f):
        return space(q, f) / cell(q, stores_row[f])[0]

    def tv(q, key):
        return T[q][key]["v"]

    def ts(qs, key):
        return "; ".join(T[q][key]["src"] for q in qs)

    def m2(x):
        return f"{x * 1000:,.0f}".replace(",", " ")

    # «Красный Яр»/«Слата»: площадь 31.12.2023 = «Пятёрочка» сопоставимая − «как отчитано»;
    # назад — чистым приростом их магазинов по кварталам × средний магазин КЯ/Слаты
    py_comp = area_by_fmt["2023H2"]["pyaterochka"]["v"]
    ky_end23 = py_comp - space("2023Q4", "pyaterochka")
    ky_n23 = tu["ky_slata_stores_2023_end"]["v"]
    ky_per_store = ky_end23 / ky_n23
    ky_h2 = tv("2023Q3", "ky_slata_net") + tv("2023Q4", "ky_slata_net")
    ky_h1 = tv("2023Q1", "ky_slata_net") + tv("2023Q2", "ky_slata_net")
    ky = {"2023H1": ky_end23 - ky_h2 * ky_per_store, "2022H2": ky_end23 - (ky_h2 + ky_h1) * ky_per_store}
    ky_net = {"2023H1": ky_h2, "2022H2": ky_h2 + ky_h1}
    ky_qs = {"2023H1": ("2023Q3", "2023Q4"), "2022H2": ("2023Q1", "2023Q2", "2023Q3", "2023Q4")}
    ky_src = f"{b24.ref(OQ, cell('2023Q4', space_row['pyaterochka'])[1])}; {tu['ky_slata_stores_2023_end']['src']}"
    A, area_est = {}, {}
    for p, q in (("2022H2", "2022Q4"), ("2023H1", "2023Q2")):
        A[p] = space(q, "pyaterochka") + space(q, "perekrestok") + space(q, "chizhik") + ky[p]
        area_est[p] = V(
            r6(A[p]), f"{refs(q, space_row.values())}; {ky_src}; {ts(ky_qs[p], 'ky_slata_net')}",
            f"{sheet}; сопоставимый базис: databook «как отчитано» на конец полугодия "
            + " + ".join(f"{name[f]} {ru(space(q, f))}" for f in fmts3)
            + f" + «Красный Яр»/«Слата» {ru(ky[p])} (их площадь 31.12.2023 {ru(ky_end23)} = «Пятёрочка» "
              f"сопоставимая {ru(py_comp)} (network.area_end_by_format[2023H2]) − «как отчитано» "
              f"{ru(space('2023Q4', 'pyaterochka'))}; минус чистый прирост их магазинов с конца полугодия до "
              f"31.12.2023 ({ky_net[p]} маг.) × средний магазин КЯ/Слаты {m2(ky_per_store)} м² "
              f"({ru(ky_end23)} / {ky_n23} маг.))")
    A["2023H2"] = area_end["2023H2"]["v"]

    C, closed_est = {}, {}
    for p, qs, q0 in (("2023H1", ("2023Q1", "2023Q2"), "2022Q4"), ("2023H2", ("2023Q3", "2023Q4"), "2023Q2")):
        n_p = sum(tv(q, "pyaterochka_closed") for q in qs)
        n_e = sum(tv(q, "perekrestok_closed") for q in qs)
        C[p] = n_p * avg_store(q0, "pyaterochka") + n_e * avg_store(q0, "perekrestok")
        closed_est[p] = V(
            r6(C[p]), f"{ts(qs, 'pyaterochka_closed')}; {ts(qs, 'perekrestok_closed')}; "
                      f"{refs(q0, (70, 62, 71, 63))}",
            f"{sheet}: закрытые магазины (трейдинг-апдейты) × средняя площадь магазина формата на "
            f"начало полугодия (databook «как отчитано»): П {n_p} маг. × {m2(avg_store(q0, 'pyaterochka'))} м² + "
            f"Пер {n_e} маг. × {m2(avg_store(q0, 'perekrestok'))} м²; «Чижик» закрытий не раскрывал — 0")
    # 2022: из трейдинг-апдейтов 2022 г. в первичке только 4 кв. (1–3 кв. нет) — закрытия каждого
    # полугодия 2022 г. = 2 × 4 кв. 2022, средний магазин на 30.06.2022 (для 2022H1 — не начало полугодия)
    n_p22, n_e22 = tv("2022Q4", "pyaterochka_closed"), tv("2022Q4", "perekrestok_closed")
    c22 = 2 * (n_p22 * avg_store("2022Q2", "pyaterochka") + n_e22 * avg_store("2022Q2", "perekrestok"))
    c22_why = {"2022H1": "трейдинг-апдейтов 1–2 кв. 2022 г. в первичке нет",
               "2022H2": "трейдинг-апдейта 3 кв. 2022 г. в первичке нет"}
    c22_date = {"2022H1": "средний магазин на 30.06.2022 (для 2022H1 — отступление от правила «на начало "
                          "полугодия»; databook «как отчитано»)",
                "2022H2": "средний магазин на 30.06.2022 — начало полугодия (databook «как отчитано»)"}
    c22_src = (f"{ts(('2022Q4',), 'pyaterochka_closed')}; {ts(('2022Q4',), 'perekrestok_closed')}; "
               f"{refs('2022Q2', (70, 62, 71, 63))}")
    gross_est = {}
    for p, q1, q0 in (("2022H1", "2022Q2", "2021Q4"), ("2022H2", "2022Q4", "2022Q2")):
        net = sum(space(q1, f) - space(q0, f) for f in fmts3)
        c22_text = (f"{sheet}: {c22_why[p]} — закрытия полугодия = 2 × закрытия 4 кв. 2022: "
                    f"2 × (П {n_p22} маг. × {m2(avg_store('2022Q2', 'pyaterochka'))} м² + Пер {n_e22} маг. × "
                    f"{m2(avg_store('2022Q2', 'perekrestok'))} м²), {c22_date[p]}; «Чижик» закрытий не "
                    f"раскрывал — 0")
        closed_est[p] = V(r6(c22), c22_src, c22_text)
        gross_est[p] = V(
            r6(net + c22), f"{refs(q1, space_row.values())}; {refs(q0, space_row.values())}",
            f"{sheet}: органический чистый прирост площади трёх форматов (databook «как отчитано», "
            f"без «Карусели» и без покупки «Красного Яра»/«Слаты») {ru(net)} + закрытая площадь {ru(c22)} "
            f"(network.closed_area_est[{p}])")
    prev = {"2023H1": ("2022H2", "area_end_est"), "2023H2": ("2023H1", "area_end_est")}
    for p in ("2023H1", "2023H2"):
        q, where = prev[p]
        here = "area_end_est" if p in area_est else "area_end"
        gross_est[p] = V(
            r6(A[p] - A[q] + C[p]),
            calc=f"{sheet}: {here}[{p}] − {where}[{q}] + closed_area_est[{p}] = {ru(A[p])} − "
                 f"{ru(A[q])} + {ru(C[p])}")
    order = ("2022H1", "2022H2", "2023H1", "2023H2")
    return (area_est, {p: gross_est[p] for p in order}, {p: closed_est[p] for p in order})


# ------------------------------------------------------------------ сборка

def build(src: dict):
    PRIMARY, MAGNIT, LENTA_DB = src["primary"], src["magnit"], src["lenta_db"]
    b26 = Book(PRIMARY / DB26, DB26)
    b24 = Book(PRIMARY / DB24, DB24)
    out = {}

    # ---------- столбцы нового databook
    c26 = {s: b26.colmap(s) for s in ("Profit and Loss", "SG&A", "EBITDA", "Cash Flow",
                                       "Financial Position", "Debt")}
    c26["Operating Results"] = b26.colmap("Operating Results", prefer_adj_row=6)
    c24 = {s: b24.colmap(s) for s in ("Profit and Loss", "SG&A", "EBITDA", "Cash Flow",
                                       "Debt", "Operating Results_Q", "Financial Position")}

    def half_new(sheet, row, period, sign=1.0, none_as_zero=False):
        """Сумма двух кварталов нового databook: (значение, источник)."""
        y, h = int(period[:4]), int(period[-1])
        qs = (1, 2) if h == 1 else (3, 4)
        cols = [c26[sheet][q_new(q, y)] for q in qs]
        vals = [b26.val(sheet, f"{c}{row}") for c in cols]
        if none_as_zero:
            vals = [v or 0.0 for v in vals]
        return sign * sum(vals), b26.ref(sheet, *[f"{c}{row}" for c in cols])

    def half_old(sheet, row, period, sign=1.0):
        y, h = int(period[:4]), int(period[-1])
        qs = (1, 2) if h == 1 else (3, 4)
        cols = [c24[sheet][q_old(q, y)] for q in qs]
        vals = [b24.val(sheet, f"{c}{row}") for c in cols]
        if any(v is None for v in vals):
            return None, b24.ref(sheet, *[f"{c}{row}" for c in cols])
        return sign * sum(vals), b24.ref(sheet, *[f"{c}{row}" for c in cols])

    def half(sheet_new, row_new, sheet_old, row_old, period, sign=1.0):
        if int(period[:4]) >= 2023:
            return half_new(sheet_new, row_new, period, sign)
        return half_old(sheet_old, row_old, period, sign)

    # ================================================================ accounting.json
    # Обесценение внеоборотных активов (кроме ППИ), нетто восстановления, млн ₽ — МСФО
    imp = {
        "2026H1": (1737 + 161 - 33,
                   f"{IFRS26} прим. 10, с. 17: убыток от обесценения ОС 1 737 + НМА 161 − восстановление 33 "
                   "(с. 11: убыток 1 898, восстановление 33; ППИ и инвест. недвижимость в 1П2026 не обесценивались)"),
        "2025H1": (1036 + 323,
                   f"{IFRS26} прим. 10, с. 17 (сравн. данные): ОС 1 036 + НМА 323; {IFRS25H1} с. 12: «убыток от обесценения 1 359»"),
        "2024H1": (1544 - 31,
                   f"{IFRS25H1} с. 12: за 6 мес. 2024 г. чистый убыток от обесценения 1 544, восстановление 31"),
        "FY2025": (5207 - 1756 + 925 + 27 - 150,
                   f"{IFRS25} прим. 10 (ОС, с. 28: убыток 5 207, восстановление 1 756), прим. 13 (НМА, с. 34: 925), "
                   "прим. 12 (инвест. недвижимость, с. 33: 27 − 150); гудвил не обесценивался (с. 25)"),
        "FY2024": (4251 - 926 + 442 + 96 - 762,
                   f"{IFRS24} прим. 10 (ОС, с. 44: 4 251 − 926), прим. 14 (НМА, с. 49: 442), "
                   "прим. 12 (инвест. недвижимость, с. 47: 96 − 762); гудвил — нет (прим. 13, с. 49)"),
        "FY2023": (1149 + 548 + 231 + 247,
                   f"{IFRS24} прим. 3, сравн. данные за 2023 г. (ОС, с. 31: чистый убыток 1 149; инвест. "
                   "недвижимость, с. 32: 231; НМА, с. 33: 548), прим. 13 (гудвил, с. 49: 247, сегмент «dark kitchen»)"),
    }
    imp_half = {
        "2026H1": (imp["2026H1"][0], imp["2026H1"][1], None),
        "2025H1": (imp["2025H1"][0], imp["2025H1"][1], None),
        "2025H2": (imp["FY2025"][0] - imp["2025H1"][0], imp["FY2025"][1] + "; " + imp["2025H1"][1],
                   "2025 г. − 1П2025"),
        "2024H2": (imp["FY2024"][0] - imp["2024H1"][0], imp["FY2024"][1] + "; " + imp["2024H1"][1],
                   "2024 г. − 1П2024"),
    }

    periods = {}
    memo = {}
    for p in ("2024H2", "2025H1", "2025H2", "2026H1"):
        rev, s_rev = half_new("Profit and Loss", 6, p)
        adj, s_adj = half_new("EBITDA", 14, p)
        lti, s_lti = half_new("EBITDA", 16, p, sign=-1.0)
        ebr, s_ebr = half_new("EBITDA", 22, p)
        oneoff, s_one = (half_new("EBITDA", 20, p, sign=-1.0, none_as_zero=True)
                         if p in ("2025H1", "2025H2", "2026H1") else (0.0, None))
        da_all, s_da = half_new("SG&A", 8, p)
        cx_ppe, s_ppe = half_new("Cash Flow", 30, p, sign=-1.0)
        cx_nma, s_nma = half_new("Cash Flow", 36, p, sign=-1.0)
        im, s_im, c_im = imp_half[p]
        da = da_all - im
        periods[p] = {
            "revenue": V(bn(rev), s_rev, "сумма кварталов полугодия; млн ₽ / 1000"),
            "adj_ebitda": V(bn(adj), s_adj, "скорр. EBITDA до МСФО 16 (без LTI и разовых), сумма кварталов"),
            "lti": V(bn(lti), s_lti, "расход LTI и прочего вознаграждения менеджмента, знак — расход положительным"),
            "ebitda_rep": V(bn(ebr), s_ebr, "EBITDA до МСФО 16 отчётная = скорр. − LTI − разовые"),
            "da": V(bn(da), s_da + "; " + s_im,
                    f"D&A и обесценение до МСФО 16 (SG&A, стр. 8) {ru(bn(da_all))} − обесценение внеоборотных "
                    f"активов нетто (МСФО) {ru(bn(im))}" + (f" ({c_im})" if c_im else "")),
            "capex": V(bn(cx_ppe + cx_nma), s_ppe + "; " + s_nma,
                       "денежный capex: приобретение ОС (с первоначальными затратами по ППИ) + НМА, до МСФО 16, "
                       "положительным"),
        }
        memo[p] = {
            "da_incl_impairment": V(bn(da_all), s_da, "D&A и обесценение до МСФО 16 (как в databook)"),
            "impairment_net": V(bn(im), s_im, c_im or "обесценение ОС + НМА (+ инвест. недвижимость) − восстановление"),
            "oneoff_ecl_adj": V(bn(oneoff) + 0.0, s_one or b26.ref("EBITDA", "(нет строки)"),
                                "разовая корректировка ОКУ (470-ФЗ), 2 кв. 2025; в скорр. EBITDA не входит")
            if s_one else V(0.0, calc="строки нет в периоде (корректировка только во 2 кв. 2025)"),
            "adj_margin": V(r6(adj / rev), calc="adj_ebitda / revenue"),
            "rep_margin": V(r6(ebr / rev), calc="ebitda_rep / revenue"),
            "capex_pct": V(r6((cx_ppe + cx_nma) / rev), calc="capex / revenue"),
            "da_pct": V(r6(da / rev), calc="da / revenue"),
        }
    out["accounting"] = {
        "as_of": AS_OF,
        "unit": "млрд ₽",
        "basis": "до МСФО 16 (IAS 17)",
        "source_note": f"{DB26} (databook X5, 2 кв. 2026): кварталы 3 кв. 2024 – 2 кв. 2026, столбцы U–AB; "
                       "обесценение — МСФО X5 (промежуточная 1П2025, 1П2026; годовая 2024, 2025)",
        "periods": periods,
        "memo": memo,
    }

    # ================================================================ network.json
    # Сопоставимый базис: «Пятёрочка» с «Красным Яром»/«Слатой», «Перекрёсток», «Чижик»; без дарксторов
    OR = "Operating Results"
    orc = c26[OR]
    TU325 = "q3_2025_trading_update_rus.pdf"
    TU425 = "q4_2025_trading_update_rus.pdf"
    TU224 = "q2_2024_trading_update_rus.pdf"
    TU423 = "q4_2023_trading_update_rus.pdf"
    area_fmt = {}   # период → формат → (тыс. м², src, calc)
    stores_fmt = {}

    def or_cell(row, key):
        col = orc[key]
        return b26.val(OR, f"{col}{row}"), b26.ref(OR, f"{col}{row}")

    # из databook (скорр. столбцы «Скорр.*»: КЯ/Слата в «Пятёрочке»)
    for p, key in (("2024H2", "4 КВ. 2024"), ("2025H1", "2 КВ. 2025"), ("2025H2", "4 КВ. 2025"),
                   ("2026H1", "2 КВ. 2026")):
        area_fmt[p] = {}
        stores_fmt[p] = {}
        for fmt, ra, rs in (("pyaterochka", 50, 45), ("perekrestok", 51, 46), ("chizhik", 52, 47)):
            va, sa = or_cell(ra, key)
            vs, ss = or_cell(rs, key)
            area_fmt[p][fmt] = (va, sa, None)
            stores_fmt[p][fmt] = (vs, ss, None)
    # 30.09.2024 (скорр.) — база для восстановления 31.12.2023 и 30.06.2024
    a_py_q3_24, s_py_q3_24 = or_cell(50, "3 КВ. 2024")
    n_py_q3_24, sn_py_q3_24 = or_cell(45, "3 КВ. 2024")
    a_pe_q3_24, s_pe_q3_24 = or_cell(51, "3 КВ. 2024")
    a_ch_q3_24, s_ch_q3_24 = or_cell(52, "3 КВ. 2024")
    n_pe_q3_24, sn_pe_q3_24 = or_cell(46, "3 КВ. 2024")
    n_ch_q3_24, sn_ch_q3_24 = or_cell(47, "3 КВ. 2024")
    tu325_ref = f"{TU325}, с. 12 «Прирост торговой площади и количества магазинов»"
    # прирост «Пятёрочки» (скорр., с КЯ/Слатой) за 3 кв. 2024 и 9М 2024 — м² и магазины
    area_fmt["2024H1"] = {
        "pyaterochka": (a_py_q3_24 - 119.342, s_py_q3_24 + "; " + tu325_ref,
                        "площадь 30.09.2024 (скорр.) − прирост за 3 кв. 2024 (119 342 м², скорр.)"),
        "perekrestok": (a_pe_q3_24 - 5.400, s_pe_q3_24 + "; " + tu325_ref,
                        "площадь 30.09.2024 − прирост за 3 кв. 2024 (5 400 м²)"),
        "chizhik": (a_ch_q3_24 - 62.966, s_ch_q3_24 + "; " + tu325_ref,
                    "площадь 30.09.2024 − прирост за 3 кв. 2024 (62 966 м²)"),
    }
    area_fmt["2023H2"] = {
        "pyaterochka": (a_py_q3_24 - 357.740, s_py_q3_24 + "; " + tu325_ref +
                        f"; сверка: {TU425}, с. 14 (прирост 12М 2024 596 044 м² от 9 177 964 м² = 8 581 920)",
                        "площадь 30.09.2024 (скорр.) − прирост за 9М 2024 (357 740 м², скорр.)"),
        "perekrestok": (a_pe_q3_24 - 14.212, s_pe_q3_24 + "; " + tu325_ref +
                        f"; сверка: {TU423}, с. 5 (31.12.2023: 1 084 913 м²)",
                        "площадь 30.09.2024 − прирост за 9М 2024 (14 212 м²)"),
        "chizhik": (a_ch_q3_24 - 146.144, s_ch_q3_24 + "; " + tu325_ref +
                    f"; сверка: {TU423}, с. 5 (31.12.2023: 442 110 м²)",
                    "площадь 30.09.2024 − прирост за 9М 2024 (146 144 м²)"),
    }
    stores_fmt["2024H1"] = {
        "pyaterochka": (n_py_q3_24 - 359, sn_py_q3_24 + "; " + tu325_ref,
                        "магазины 30.09.2024 (скорр.) − чистый прирост за 3 кв. 2024 (359)"),
        "perekrestok": (n_pe_q3_24 - 5, sn_pe_q3_24 + "; " + tu325_ref, "30.09.2024 − 5"),
        "chizhik": (n_ch_q3_24 - 217, sn_ch_q3_24 + "; " + tu325_ref, "30.09.2024 − 217"),
    }
    stores_fmt["2023H2"] = {
        "pyaterochka": (n_py_q3_24 - 1050, sn_py_q3_24 + "; " + tu325_ref,
                        "магазины 30.09.2024 (скорр.) − чистый прирост за 9М 2024 (1 050)"),
        "perekrestok": (n_pe_q3_24 - 20, sn_pe_q3_24 + "; " + tu325_ref, "30.09.2024 − 20"),
        "chizhik": (n_ch_q3_24 - 502, sn_ch_q3_24 + "; " + tu325_ref, "30.09.2024 − 502"),
    }
    order = ["2023H2", "2024H1", "2024H2", "2025H1", "2025H2", "2026H1"]
    fmts = ("pyaterochka", "perekrestok", "chizhik")
    area_end, area_by_fmt, stores_by_fmt, stores_end_all = {}, {}, {}, {}
    for p in order:
        tot = sum(area_fmt[p][f][0] for f in fmts)
        srcs = "; ".join(dict.fromkeys(area_fmt[p][f][1] for f in fmts))
        area_end[p] = V(r6(tot), srcs, "«Пятёрочка» (с «Красным Яром» и «Слатой») + «Перекрёсток» + «Чижик», "
                                       "без дарксторов и «тёмных» кухонь")
        area_by_fmt[p] = {f: V(r6(area_fmt[p][f][0]), area_fmt[p][f][1], area_fmt[p][f][2]) for f in fmts}
        stores_by_fmt[p] = {f: V(stores_fmt[p][f][0], stores_fmt[p][f][1], stores_fmt[p][f][2]) for f in fmts}
        stores_end_all[p] = V(sum(stores_fmt[p][f][0] for f in fmts), calc="сумма трёх форматов (сопоставимый базис)")
    # сверки с раскрытием итогов
    checks = {
        "2024H2": (10961.011, f"{PR425}, с. 4 (итого 31.12.2024 без дарксторов, исторические данные скорректированы)"),
        "2025H1": (11376.156, f"{PR226}, с. 4 (итого 30.06.2025)"),
        "2025H2": (11912.101, f"{PR226}, с. 4 (итого 31.12.2025)"),
        "2026H1": (12157.240, f"{PR226}, с. 4 (итого 30.06.2026)"),
    }
    for p, (v, s) in checks.items():
        assert abs(area_end[p]["v"] - v) < 0.002, (p, area_end[p]["v"], v)
        area_end[p]["check"] = s

    net_added = {}
    for i in range(1, len(order)):
        p, q = order[i], order[i - 1]
        net_added[p] = V(r6(area_end[p]["v"] - area_end[q]["v"]), calc=f"area_end[{p}] − area_end[{q}]")

    # Открытия и закрытия магазинов по кварталам (трейдинг-апдейты X5), шт.
    # opened — «без учёта закрытий»; closed — раскрыто прямо или = opened − net; None — не раскрыто.
    tu = {
        "2024Q1": ("q1-2024_trading_update_rus.pdf", "с. 1, 4, 6",
                   {"pyaterochka": (399 + 4, 116, 287), "perekrestok": (12, 6, 6), "chizhik": (116, 0, 116)},
                   "«Пятёрочка» 399 открытий, «Красный Яр»/«Слата» +2 чистыми при 2 закрытых «ХлебСоль» (группа КЯ) → "
                   "в сопоставимом базисе открыто 403, закрыто 116 (114 «у дома» + 2 «ХлебСоль»), чистый +287; "
                   "«Перекрёсток» +6 чистыми при 6 закрытых → 12 открытий; «Чижик» открыл 116, закрытий нет"),
        "2024Q2": ("q2_2024_trading_update_rus.pdf", "с. 1, 4, 6",
                   {"pyaterochka": (470, 66, 404), "perekrestok": (14, 5, 9), "chizhik": (170, 1, 169)},
                   "открыто без учёта закрытий: «Пятёрочка» 470, «Чижик» 170; закрыто 66 / 5 / 1; "
                   "чистые 404 / 9 / 169 (открытия «Перекрёстка» = чистые + закрытые)"),
        "2024Q3": ("q3-2024_trading_update_rus.pdf", "с. 1, 4, 6",
                   {"pyaterochka": (448, 89, 359), "perekrestok": (10, 5, 5), "chizhik": (219, 2, 217)},
                   "открыто 448 / — / 219; закрыто 89 / 5 / 2; чистые 359 / 5 / 217"),
        "2024Q4": ("q4_2024_trading_update_rus.pdf", "с. 6–8, 15",
                   {"pyaterochka": (722, 102, 620), "perekrestok": (1, 7, -6), "chizhik": (347, 3, 344)},
                   "открыто 721 / — / 347; закрыто 102 / 7 / 3; чистые 619 / −6 / 344 (с. 15); в сопоставимом "
                   "базисе чистый прирост «Пятёрочки» за 4 кв. 2024 — 620 (q4_2025_trading_update_rus.pdf, с. 14, "
                   "скорр. с КЯ/Слатой) → +1 магазин КЯ/Слаты учтён как открытие: 722"),
        "2025Q1": ("q1_2025_trading_update_rus.pdf", "с. 1, 4–5, 13",
                   {"pyaterochka": (440, 54, 386), "perekrestok": (5, 7, -2), "chizhik": (129, 4, 125)},
                   "открыто 440 / 5 / 129; закрыто 53 «у дома» + 1 «Слата» / 7 / 4; "
                   "чистые (сопоставимые, databook «Скорр.») 386 / −2 / 125"),
        "2025Q2": ("q2_2025_trading_update_rus.pdf", "с. 1, 4–5, 13",
                   {"pyaterochka": (613, 73, 540), "perekrestok": (17, 9, 8), "chizhik": (210, 5, 205)},
                   "открыто 613 / 17 / 210; закрыто 65 «у дома» + 8 «Красный Яр»/«Слата» (интеграция в «Пятёрочку») "
                   "/ 9 / 5; чистые в сопоставимом базисе 540 / 8 / 205"),
        "2025Q3": (TU325, "с. 3–4, 12",
                   {"pyaterochka": (644, 168, 476), "perekrestok": (12, 6, 6), "chizhik": (237, 4, 233)},
                   "открыто без учёта закрытий 644 / 12 / 237; закрытия не названы — = открытия − чистые (476 / 6 / 233)"),
        "2025Q4": (TU425, "с. 5–6, 14",
                   {"pyaterochka": (672, 129, 543), "perekrestok": (11, 5, 6), "chizhik": (None, None, 344)},
                   "открыто в современной концепции «без учёта закрытий»: «Пятёрочка» 672, «Перекрёсток» 11; "
                   "закрытия = открытия − чистые (543 / 6); «Чижик» — только чистые +344"),
        "2026Q1": ("q1_2026_trading_update_rus.pdf", "с. 3–4, 12",
                   {"pyaterochka": (381, 164, 217), "perekrestok": (3, 11, -8), "chizhik": (None, None, 139)},
                   "открыто «без учёта закрытий»: «Пятёрочка» 381, «Перекрёсток» 3; чистые 217 / −8 / 139; "
                   "закрытия = открытия − чистые; «Чижик» — только чистые"),
        "2026Q2": ("q2_2026_trading_update_rus.pdf", "с. 3–4, 12",
                   {"pyaterochka": (450, 196, 254), "perekrestok": (11, 3, 8), "chizhik": (None, None, 204)},
                   "открыто «без учёта закрытий»: «Пятёрочка» 450, «Перекрёсток» 11; чистые 254 / 8 / 204; "
                   "закрытия = открытия − чистые; «Чижик» — только чистые"),
    }
    oc = {}
    for qk, (doc, pages, data, how) in tu.items():
        row = {}
        for f in fmts:
            o, c, n = data[f]
            s = f"{doc}, {pages}"
            row[f] = {
                "opened": V(o, s, how) if o is not None else V(None, calc="не раскрыто"),
                "closed": V(c, s, how) if c is not None else V(None, calc="не раскрыто"),
                "net": V(n, s, how),
            }
        oc[qk] = row
    # проверка: чистые по кварталам полугодия = разности магазинов на концах полугодий
    half_q = {"2024H1": ("2024Q1", "2024Q2"), "2024H2": ("2024Q3", "2024Q4"),
              "2025H1": ("2025Q1", "2025Q2"), "2025H2": ("2025Q3", "2025Q4"),
              "2026H1": ("2026Q1", "2026Q2")}
    stores_check = {}
    for i in range(1, len(order)):
        p, q = order[i], order[i - 1]
        for f in fmts:
            net_q = sum(tu[k][2][f][2] for k in half_q[p])
            net_b = stores_fmt[p][f][0] - stores_fmt[q][f][0]
            stores_check[f"{p}.{f}"] = (net_q, net_b)
    # закрытая площадь (оценка) и валовые открытия площади
    closed_area, gross = {}, {}
    for i in range(1, len(order)):
        p, q = order[i], order[i - 1]
        parts, total_closed = [], 0.0
        for f in fmts:
            avg = area_fmt[q][f][0] / stores_fmt[q][f][0]   # тыс. м² на магазин на начало полугодия
            cl = [tu[k][2][f][1] for k in half_q[p]]
            n_closed = sum(c for c in cl if c is not None)
            unknown = any(c is None for c in cl)
            total_closed += n_closed * avg
            parts.append(f"{f}: {n_closed} маг. × {avg * 1000:.0f} м²" + (" (закрытия части кварталов не раскрыты — в них 0)" if unknown else ""))
        closed_area[p] = V(r6(total_closed), calc="закрытые магазины (трейдинг-апдейты) × средняя площадь магазина "
                                                  "формата на начало полугодия; " + "; ".join(parts))
        gross[p] = V(r6(net_added[p]["v"] + total_closed),
                     calc=f"оценка: чистый прирост площади {ru(net_added[p]['v'])} + закрытая площадь "
                          f"{ru(total_closed)} (network.closed_area_est[{p}])")
    close_rate = {}
    for i in range(1, len(order)):
        p, q = order[i], order[i - 1]
        close_rate[p] = V(r6(closed_area[p]["v"] / area_end[q]["v"] * 2),
                          calc=f"closed_area_est[{p}] / area_end[{q}] × 2 (годовая доля закрываемой площади)")
    # история индекса эффективной площади до 31.12.2023 (docs/MODEL.md §4.1) — оценки листа сети
    area_est, gross_early, closed_early = network_history_est(b24, c24["Operating Results_Q"], area_end,
                                                              area_by_fmt)
    gross = {**gross_early, **gross}
    closed_area = {**closed_early, **closed_area}

    by_fmt = {f: {"area": area_by_fmt["2026H1"][f], "stores": stores_by_fmt["2026H1"][f]} for f in fmts}
    out["network"] = {
        "as_of": AS_OF,
        "unit": "площадь — тыс. м²; магазины — шт.",
        "basis": "сопоставимый: «Пятёрочка» с «Красным Яром» и «Слатой», «Перекрёсток», «Чижик»; без дарксторов "
                 "Vprok.ru, совместных дарксторов и «тёмных» кухонь «Много лосося» (они были в итоге databook "
                 "со 2 кв. 2024 по 3 кв. 2025)",
        "area_end": area_end,
        "area_end_est": area_est,
        "net_added": net_added,
        "gross_opened_hist": gross,
        "closed_area_est": closed_area,
        "close_rate_implied": close_rate,
        "stores_end": {"2026H1": V(30604, b26.ref(OR, f"{orc['2 КВ. 2026']}48"),
                                   "итого магазинов 30.06.2026 (= сумма трёх форматов)")},
        "stores_end_hist": stores_end_all,
        "area_end_by_format": area_by_fmt,
        "stores_end_by_format": stores_by_fmt,
        "by_format_2026H1": by_fmt,
        "openings_closures_stores": oc,
    }
    out["_stores_check"] = stores_check

    # ================================================================ balance.json
    FPs = "Financial Position"
    R = c26[FPs]["2026-06-30"]
    D = c26["Debt"]["2026-06-30"]

    def fp(row):
        return b26.val(FPs, f"{R}{row}"), b26.ref(FPs, f"{R}{row}")

    def debt(row):
        return b26.val("Debt", f"{D}{row}"), b26.ref("Debt", f"{D}{row}")

    td, s_td = debt(6)
    nd, s_nd = debt(11)
    lev, s_lev = debt(12)
    lease, s_lease = debt(13)
    cash, s_cash = fp(26)
    lt_b, s_ltb = fp(42)
    st_b, s_stb = fp(49)
    inv, s_inv = fp(19)
    rec, s_rec = fp(23)
    tp, s_tp = fp(48)
    itp, s_itp = fp(53)
    itr, s_itr = fp(24)
    cl_st, s_cl = fp(52)
    vat, s_vat = fp(25)
    assert abs(td - cash - nd) < 1.0
    borrowings = lt_b + st_b
    nwc_trade = inv + rec - tp
    other_tax = 53115.0
    s_other_tax = f"{IFRS26} прим. 13, с. 19: «Налоги, кроме налога на прибыль» (краткосрочные) 53 115"
    nwc_ml = nwc_trade - other_tax - itp - cl_st
    # операционная часть «Резервов и прочих краткосрочных обязательств» (МСФО 1П2026, прим. 13, с. 19);
    # не входят: дивиденды к выплате, кредиторка за ОС/НМА/бизнесы, налоговые резервы, пут на НДУ
    n13 = f"{IFRS26} прим. 13, с. 19: "
    op_liab = {
        "other_payables_accruals": (51449.0, n13 + "«Прочая кредиторская задолженность и начисления» 51 449"),
        "personnel": (44994.0, n13 + "«Обязательства перед персоналом» 44 994"),
        "other_taxes": (other_tax, n13 + "«Налоги, кроме налога на прибыль» 53 115"),
        "advances_received": (2982.0, n13 + "«Авансы полученные» 2 982"),
        "other_nonfinancial": (3915.0, n13 + "«Прочие краткосрочные нефинансовые обязательства» 3 915"),
    }
    nwc_op = nwc_trade + vat - cl_st - sum(x for x, _ in op_liab.values())
    rev_ltm = periods["2025H2"]["revenue"]["v"] + periods["2026H1"]["revenue"]["v"]
    rep_ltm = periods["2025H2"]["ebitda_rep"]["v"] + periods["2026H1"]["ebitda_rep"]["v"]
    adj_ltm = periods["2025H2"]["adj_ebitda"]["v"] + periods["2026H1"]["adj_ebitda"]["v"]
    out["balance"] = {
        "as_of": AS_OF,
        "unit": "млрд ₽",
        "basis": "до МСФО 16 (IAS 17); аренда МСФО 16 — справочно, вне долга",
        "net_debt": V(bn(nd), s_nd + f"; {PR226}, с. 12 (чистый долг 310 647, 1,08х)",
                      "чистый долг компании до МСФО 16 = общий долг (с лизингом) − денежные средства; "
                      "краткосрочные финансовые вложения не вычитаются"),
        "total_debt": V(bn(td), s_td + f"; {PR226}, с. 12, сноска 18 («включает обязательства по лизингу»)"),
        "cash": V(bn(cash), s_cash + f"; {IFRS26} прим. 8, с. 15"),
        "borrowings": V(bn(borrowings), b26.ref(FPs, f"{R}42", f"{R}49") + f"; {IFRS26} прим. 15, с. 20 (434 709)",
                        "кредиты и займы долгосрочные + краткосрочные (баланс)"),
        "leasing": V(bn(td - borrowings), s_td + "; " + b26.ref(FPs, f"{R}42", f"{R}49"),
                     "общий долг компании − кредиты и займы баланса (лизинг внутри «общего долга», отдельно не раскрыт)"),
        "accrued_interest": V(bn(fp(51)[0]), fp(51)[1], "в «общий долг» компании не входит — строка моста"),
        "dividends_payable": V(60.425, f"{IFRS26} прим. 13, с. 19: «Дивиденды к выплате 60 425»",
                               "объявленные и не выплаченные на 30.06.2026: финальные за 2025 г. 60,265 + "
                               "невостребованные прошлых выплат 0,160"),
        "nwc_trade": V(bn(nwc_trade), s_inv + "; " + s_rec + "; " + s_tp,
                       "запасы + торговая и прочая дебиторская задолженность с авансами выданными − торговая "
                       "кредиторская задолженность"),
        "nwc_magnit_like": V(bn(nwc_ml), s_inv + "; " + s_rec + "; " + s_tp + "; " + s_itp + "; " + s_cl + "; " +
                             s_other_tax,
                             "nwc_trade − налоги, кроме налога на прибыль (краткоср.) − налог на прибыль к уплате − "
                             "краткосрочные обязательства по договорам (определение книги Магнита, research/X1 §5.1)"),
        "nwc_op": V(bn(nwc_op), calc="запасы + торговая и прочая дебиторка с авансами + НДС и прочие налоги к "
                                     "возмещению − торговая кредиторка − краткосрочные обязательства по договорам − "
                                     "операционная часть «Резервов и прочих обязательств» (прочая кредиторка и "
                                     "начисления, обязательства перед персоналом, налоги кроме налога на прибыль, "
                                     "авансы полученные, прочие нефинансовые); не входят дивиденды к выплате, "
                                     "кредиторка за ОС/НМА/бизнесы, налоговые резервы, выкуп НДУ (строки моста или "
                                     "capex), налог на прибыль к уплате и к возмещению (строка моста income_tax_net). "
                                     "Определение книги A-W1 (data/assumptions/evidence/capex-wc-tax/README.md §9)"),
        "nwc": V(bn(nwc_op), calc="= nwc_op (определение книги working_capital, A-W1)"),
        "nwc_components": {
            "inventories": V(bn(inv), s_inv),
            "receivables_and_advances": V(bn(rec), s_rec),
            "trade_payables": V(bn(tp), s_tp),
            "other_taxes_payable": V(bn(other_tax), s_other_tax),
            "income_tax_payable": V(bn(itp), s_itp),
            "contract_liabilities_st": V(bn(cl_st), s_cl),
            "vat_other_taxes_receivable": V(bn(vat), s_vat),
            "operating_other_liabilities": {k: V(bn(x), s) for k, (x, s) in op_liab.items()},
        },
        "revenue_ltm": V(r6(rev_ltm), calc="accounting 2025H2 + 2026H1 (= 2025 − 1П2025 + 1П2026)"),
        "ebitda_rep_ltm": V(r6(rep_ltm), calc="accounting 2025H2 + 2026H1, EBITDA до МСФО 16 отчётная"),
        "adj_ebitda_ltm": V(r6(adj_ltm), calc="accounting 2025H2 + 2026H1, скорр. EBITDA до МСФО 16"),
        "net_debt_to_ebitda": V(r6(lev), s_lev + f"; {PR226}, с. 12; {IFRS26} прим. 16, с. 21 (1,08x)",
                                "чистый долг / EBITDA до МСФО 16 LTM (ковенант ≤ 4,00х)"),
        "lease_liabilities_ifrs16": V(bn(lease), s_lease + f"; {IFRS26} прим. 11, с. 18"),
        "credit_lines_unused": V(804.643, f"{IFRS26} прим. 24, с. 25; {PR226}, с. 12",
                                 "невыбранные лимиты кредитных линий крупнейших банков"),
        "st_investments": V(bn(fp(22)[0]), fp(22)[1]),
    }

    # ================================================================ bridge.json
    tax_prov = 10167.0
    comp_asset, s_comp = fp(20)
    assoc, s_assoc = fp(12)
    sti, s_sti = fp(22)
    out["bridge"] = {
        "as_of": AS_OF,
        "unit": "млрд ₽",
        "sign": "как у требования: обязательство положительное, актив отрицательный",
        "lines": [
            {"key": "accrued_interest", "label": "Начисленные проценты",
             "amount": V(bn(fp(51)[0]), fp(51)[1], "не входят в «общий долг» и чистый долг компании")},
            {"key": "nci_put", "label": "Пут на неконтролирующую долю",
             "amount": V(0.698, f"{IFRS26} прим. 13, с. 19: «Обязательства по покупке неконтролирующих долей "
                                "участия» 698 (долгосрочные; краткосрочных — «–»)")},
            {"key": "lti_liability", "label": "Обязательство LTI",
             "amount": V(7.896, f"{IFRS26} прим. 13, с. 19: «Обязательства по долгосрочным программам премирования "
                                "ключевых сотрудников» 7 896",
                         "только долгосрочная часть; краткосрочная внутри «обязательств перед персоналом» 44 994 "
                         "отдельно не раскрыта")},
            {"key": "tax_provisions_net", "label": "Резервы по налоговым позициям (нетто)",
             "amount": V(bn(tax_prov - comp_asset),
                         f"{IFRS26} прим. 13, с. 19: «Резервы и обязательства по неопределённостям в отношении правил "
                         f"исчисления налогов» 10 167; {s_comp} («Компенсирующий актив»)",
                         "10 167 − компенсирующий актив 2 568 (возмещения продавцов купленных бизнесов)")},
            {"key": "income_tax_net", "label": "Налог на прибыль к уплате − к возмещению (нетто)",
             "amount": V(bn(itp - itr), b26.ref(FPs, f"{R}53") + " − " + f"{R}24",
                         f"налог на прибыль к уплате {ru(bn(itp))} − к возмещению {ru(bn(itr))}; вне NWC (налог "
                         "модели денежный в периоде, A-W1) — остаток якоря гасится в 2П2026; знак минус — "
                         "требование к бюджету")},
            {"key": "deferred_consideration", "label": "Отложенное возмещение по приобретениям бизнесов",
             "amount": V(1.244, f"{IFRS26} прим. 6, с. 14: «Переданное за отчетный период возмещение включало "
                                "денежное возмещение в размере 23 млн руб. и отложенное возмещение в размере "
                                "1 244 млн руб.»",
                         "нижняя граница: отложенное возмещение по сделкам 1П2026; остаток на 30.06 отдельно не "
                         "раскрыт (внутри кредиторки за ОС, НМА и бизнесы 24 917; за сделки прошлых периодов в "
                         "1П2026 перечислено 40 млн); купленные бизнесы уже в якоре, платёж — ни в NWC, ни в capex")},
            {"key": "st_investments", "label": "Краткосрочные финансовые вложения",
             "amount": V(-bn(sti), s_sti, "актив — со знаком минус")},
            {"key": "associates", "label": "Инвестиции в ассоциированные и СП",
             "amount": V(-bn(assoc), s_assoc,
                         "балансовая стоимость; актив — со знаком минус")},
        ],
        "outside_bridge": {
            "lease_liabilities_ifrs16": V(bn(lease), s_lease, "базис до МСФО 16: аренда внутри EBITDA, вне моста"),
            "capital_commitments": V(29.882, f"{IFRS26} прим. 26, с. 26", "обязательства по капвложениям — "
                                                                          "это capex прогноза, не мост"),
            "treasury_shares_mln": V(25.591692, f"{IFRS26} прим. 17, с. 21",
                                     "не актив моста; в формуле цены — выручка от продажи пакета n·k·P_рынок, "
                                     "знаменатель — выпущенные акции N + n (docs/MODEL.md §7.2; ядро читает n "
                                     "из shares.treasury)"),
            "gorod77_consideration": V(None, calc="цена «Города 77» (20.07.2026) не раскрыта; после отчётной даты"),
            "obligation_to_x5_retail_group_nv": V(0.0, f"{IFRS26} прим. 17, с. 22; ОДДС с. 7 (выплачено 4 638 во "
                                                       "2 кв. 2026)", "обязательств перед бывшей материнской нет"),
        },
    }

    # ================================================================ shares.json
    out["shares"] = {
        "as_of": AS_OF,
        "unit": "млн шт.",
        "issued": V(271.572872, f"{IFRS26} прим. 17, с. 21 (271 572 872 обыкновенных акций); "
                                "moex-iss-x5-share-tqbr-2026-09-28.json › ISSUESIZE", "шт. / 10⁶"),
        "treasury": V(25.591692, f"{IFRS26} прим. 17, с. 21 (25 591 692 собственных выкупленных, 54 287 млн ₽)",
                      "шт. / 10⁶; включают 15 770 акций у дочерних обществ"),
        "treasury_at_subsidiaries": V(0.01577, f"{IFRS26} прим. 17, с. 21 (15 770 акций у дочерних организаций); "
                                               f"{ER6M26}, с. 38", "шт. / 10⁶"),
        "outstanding_mln": V(245.98118, f"{IFRS26} прим. 17, с. 21 (245 981 180 в обращении)",
                             "выпущено − казначейские = 271 572 872 − 25 591 692"),
        "free_float_pct": V(0.29, "x5-investors-shares-2026-09-28.html (x5.ru/ru/investors/shares/, доступ "
                                  "28.09.2026): «free-float … по методике ПАО Московская Биржа … 29%»",
                            "доля единицы"),
        "treasury_book_value": V(54.287, b26.ref(FPs, f"{R}38") + f"; {IFRS26} прим. 17", "млрд ₽"),
        "dividend_base_mln": V(245.99695, "x5-investors-dividends-2026-09-28.html: 60 269 252 750 ₽ / 245 ₽",
                               "акции, на которые начислен финальный дивиденд за 2025 г. = в обращении + 15 770 "
                               "акций «дочек» (внутригрупповые)"),
        "weighted_avg_eps_2026H1": V(245.272401, f"{IFRS26} прим. 18, с. 22"),
        "treasury_sales_history": [
            {"period": "2026H1",
             "units_mln": V(0.861, f"{IFRS26} прим. 17, с. 21 («реализовала 861 000 собственных акций»)"),
             "proceeds": V(2.010, b26.ref("Cash Flow", f"{c26['Cash Flow']['2 КВ. 2026']}55") +
                           " («Поступления от реализации собственных акций» 2 010)", "млрд ₽"),
             "avg_price_rub": V(round(2010.0 / 0.861, 1), calc="2 010 млн ₽ / 0,861 млн шт."),
             "buyer": "не раскрыт"},
        ],
        "treasury_policy": {
            "text": "Наблюдательный совет 13.11.2025 одобрил отчуждение казначейских акций (9,7 % УК) в срок не "
                    "более трёх лет; погашение не объявлялось",
            "src": "x5-news-2025-11-13-dividends-9m2025-treasury.html; x5-ar25.pdf, с. 225",
        },
    }

    # ================================================================ dividends.json
    ER = ER6M26
    out["dividends"] = {
        "unit": "млрд ₽; DPS — ₽ на акцию",
        "policy": {
            "target_leverage": [1.2, 1.4],
            "no_pay_above": 2.0,
            "frequency": "дважды в год: за предыдущий год и за 9 месяцев текущего",
            "base": "свободный денежный поток при целевом значении чистый долг / EBITDA до МСФО 16 = 1,2–1,4× на конец "
                    "года, в котором планируется выплата; при текущем или прогнозном показателе выше 2,0× дивиденды "
                    "не выплачиваются",
            "approved": "Наблюдательный совет 20.03.2025, на четыре года",
            "src": "x5-ar25.pdf, с. 224 (раздел «Дивиденды»); x5-investor-presentation_rus.pdf, с. 27; "
                   "x5_q4_2024_financial_results_rus.pdf, с. 2",
        },
        "register": [
            {"id": "FY2025-final", "label": "Финальный за 2025 г.", "dps": 245.0,
             "amount": V(60.265, f"{IFRS26} прим. 17, с. 22: «60 265 млн руб. (245 руб. на акцию без учёта "
                                 "собственных выкупленных акций)»",
                         "начислено на 60,269 млрд (с внутригрупповыми 15 770 акциями «дочек»), внешним — 60,265"),
             "decided_on": "2026-06-26", "record_date": "2026-07-07", "ex_date": "2026-07-07",
             "last_cum_date": "2026-07-06",
             "pay_until": "2026-08-11", "pay_until_nominee": "2026-07-21",
             "status": "paid",
             "paid_share_on": {"2026-08-13": V(0.761119, f"{ER}, с. 42: выплачено 45 872 096 760 ₽ "
                                                        "(76,1119387862 %) на 13.08.2026")},
             "cash_in_company_at_facts_date": True,
             "src": f"x5-investors-dividends-2026-09-28.html (реестр и экс-дата 07.07.2026, последний день с "
                    f"дивидендом 06.07.2026); {ER}, с. 40–41 (ГОСА 26.06.2026, реестр 07.07.2026, выплата "
                    "номинальным держателям до 21.07.2026, прочим — до 11.08.2026)",
             "note": "ex_date — первый день торгов без права на дивиденд (T+1: день закрытия реестра). На 30.06.2026 "
                     "объявлен и не выплачен, деньги в кассе. Срок выплаты истёк 11.08.2026 (status на дату "
                     "сбора 28.09.2026 — paid); на 13.08.2026 выплачено 76,1 %, остаток 14,4 млрд не востребован "
                     "(нет реквизитов, возврат депозитариями) и остаётся обязательством."},
            {"id": "unclaimed-old", "label": "Невостребованные дивиденды прошлых выплат (за 2024 г. и 9М 2025)",
             "dps": None,
             "amount": V(0.160, f"{IFRS26} прим. 13, с. 19 (дивиденды к выплате 60 425); прим. 17, с. 22 (финальные "
                                f"60 265); {ER}, с. 42 (невыплачено: за 2024 г. 158 848 095 600 − 158 747 171 317 = "
                                "0,101 млрд; за 9М 2025 90 210 029 600 − 90 151 100 159 = 0,059 млрд)",
                         "60,425 − 60,265 (≈ 0,101 + 0,059)"),
             "decided_on": "2025-12-18", "record_date": "2026-01-06", "ex_date": "2026-01-06",
             "status": "unclaimed",
             "cash_in_company_at_facts_date": True,
             "src": "x5-investors-dividends-2026-09-28.html (реестры 09.07.2025 и 06.01.2026)",
             "note": "ex_date — последняя из двух прошедших экс-дат (09.07.2025, 06.01.2026): требование действует "
                     "на любую дату оценки после якоря, пока остаток не востребован"},
        ],
        "history": [
            {"period": "FY2024", "label": "за 2024 г.", "dps": 648.0,
             "amount": V(158.848, "x5-investors-dividends-2026-09-28.html (158 848 095 600 ₽); "
                                  f"{IFRS26} прим. 17, с. 22"),
             "decided_on": "2025-06-27", "record_date": "2025-07-09", "ex_date": "2025-07-09",
             "paid_share": V(0.999365, f"{ER}, с. 42 (99,9364649084 %)"), "payer": "ПАО «КЦ ИКС 5»"},
            {"period": "9M2025", "label": "за 9 мес. 2025 г.", "dps": 368.0,
             "amount": V(90.210, "x5-investors-dividends-2026-09-28.html (90 210 029 600 ₽)"),
             "decided_on": "2025-12-18", "record_date": "2026-01-06", "ex_date": "2026-01-06",
             "paid_share": V(0.999347, f"{ER}, с. 42 (99,9346752891 %)"), "payer": "ПАО «КЦ ИКС 5»"},
            {"period": "FY2025", "label": "финальный за 2025 г.", "dps": 245.0,
             "amount": V(60.269, "x5-investors-dividends-2026-09-28.html (60 269 252 750 ₽)"),
             "decided_on": "2026-06-26", "record_date": "2026-07-07", "ex_date": "2026-07-07",
             "paid_share": V(0.761119, f"{ER}, с. 42 (76,1119387862 % на 13.08.2026)"), "payer": "ПАО «КЦ ИКС 5»"},
        ] + [
            {"period": per, "label": lab, "dps": dps, "unit_dps": "₽ на ГДР",
             "amount": V(amt, "x5-ar25.pdf, с. 224 (график «Дивидендная история», ГДР X5 Retail Group N.V.)"),
             "record_date": None, "payer": "X5 Retail Group N.V. (ГДР)",
             "note": "выплата лица-предшественника по ГДР; дата реестра в собранных источниках не раскрыта"}
            for per, lab, amt, dps in (
                ("FY2017", "за 2017 г. (выплата 2018 г.)", 21.590, 79.5),
                ("FY2018", "за 2018 г. (выплата 2019 г.)", 25.000, 92.06),
                ("FY2019", "за 2019 г. (выплата 2020 г.)", 30.000, 110.47),
                ("9M2020", "за 9 мес. 2020 г. (выплата 2020 г.)", 19.997, 73.645),
                ("FY2020", "за 2020 г. (выплата 2021 г.)", 30.000, 110.49),
                ("9M2021", "за 9 мес. 2021 г. (выплата 2021 г.)", 20.000, 73.65))
        ],
        "history_note": "За 2021 г. (финал), 2022 и 2023 гг. дивиденды не выплачивались (x5-ar25.pdf, с. 224; "
                        "МСФО 2025 прим. 22: «в 2024 году дивиденды не объявлялись и не выплачивались»). "
                        "Выплаты N.V. по годам (ОДДС старого databook): 21,6 (2018), 25,0 (2019), 50,0 (2020), 50,0 (2021).",
        "next_expected": {
            "label": "за 9 мес. 2026 г.",
            "status": "не объявлен на 28.09.2026",
            "board_est": "≈ середина ноября 2026 (прецедент: 13.11.2025)",
            "record_date_est": "≈ начало января 2027 (прецедент: 06.01.2026)",
            "src": "research/indicators/x5_calendar.json (прецедент 2025 г.)",
        },
    }
    reg_sum = sum(r["amount"]["v"] for r in out["dividends"]["register"] if r["cash_in_company_at_facts_date"])
    assert abs(reg_sum - 60.425) < 1e-9

    # ================================================================ debt_register.json
    rd = json.loads(src["debt_research"].read_text(encoding="utf-8"))
    iss_bonds = json.loads((PRIMARY / "moex-iss-x5-bonds-tqcb-2026-09-28.json").read_text(encoding="utf-8"))

    def tab(block):
        return {r[0]: dict(zip(block["columns"], r)) for r in block["data"]}

    md = tab(iss_bonds["marketdata"])
    sec = tab(iss_bonds["securities"])
    ISSB = "moex-iss-x5-bonds-tqcb-2026-09-28.json"
    bonds = []
    for b in rd["bonds"]:
        isin = b["isin"]
        m = md.get(b["secid"], {})
        s = sec.get(b["secid"], {})
        c = b.get("coupon", {})
        mk = b.get("market_2026_09_28", {})
        live = b["status"] == "live"
        put_left = b["series"] == "003P-07"   # частично выкуплен на оферте: остаток ≠ объёму выпуска
        bonds.append({
            "series": b["series"], "isin": isin, "name": s.get("SHORTNAME") or b["series"],
            "issuer": b["issuer"],
            "outstanding_2026_06_30": V(b["outstanding_2026_06_30_rub_bn"],
                                        f"{IFRS26} прим. 15, с. 20 (балансовая {b['carrying_2026_06_30_rub_bn']}); "
                                        + ("остаток после оферты 25.09.2025 — по балансу МСФО (число бумаг не "
                                           "раскрыто; объём выпуска в ISS 21,0 — до оферты)" if put_left else
                                           "номинал = объём выпуска (ISS ISSUESIZEPLACED)"))
            if b["outstanding_2026_06_30_rub_bn"] is not None else V(None, calc="размещён после 30.06.2026"),
            "outstanding_2026_09_28": V(b["outstanding_2026_09_28_rub_bn"],
                                        (f"{IFRS26} прим. 15, с. 20: остаток после оферты 25.09.2025 по балансу "
                                         "МСФО; новых оферт до 28.09.2026 не было (следующая — 19.07.2027)"
                                         if put_left else
                                         f"{ISSB} › securities.ISSUESIZEPLACED; пресс-релизы о размещениях 003P-20/21")
                                        if live else f"{ISSB}; торги 003P-04 прекращены 07.09.2026 (оферта 11.09.2026)"),
            "carrying_2026_06_30": V(b["carrying_2026_06_30_rub_bn"], f"{IFRS26} прим. 15, с. 20")
            if b["carrying_2026_06_30_rub_bn"] is not None else V(None, calc="размещён после 30.06.2026"),
            "coupon_type": b["coupon_type"],
            "coupon_formula": c.get("formula"),
            "spread_to_key_rate": V(round(c.get("spread_pp") / 100.0, 6), b["source"]) if c.get("spread_pp") is not None
            else None,
            "coupon_fixed": V(round(c.get("rate_pct") / 100.0, 6), b["source"]) if c.get("rate_pct") is not None else None,
            "coupon_now": (V(round(s["COUPONPERCENT"] / 100.0, 6), f"{ISSB} › securities.COUPONPERCENT "
                                                                     "(текущий купон)")
                           if live and s.get("COUPONPERCENT") else
                           V(round(b["coupon_now_pct_est"] / 100.0, 6),
                             calc="оценка: ключевая ставка 14,00 % + спред выпуска (research/X2 §3.2)")
                           if live and b.get("coupon_now_pct_est") else V(None, calc="выпуск выбыл")),
            "put_date": b.get("next_put_or_offer"),
            "maturity": b.get("legal_maturity"),
            "repayment_date_for_model": b.get("repayment_date_for_model"),
            "price": V(m.get("LAST"), f"{ISSB} › marketdata.LAST (28.09.2026 ≈11:44 МСК)") if m.get("LAST") else None,
            "ytm": V(round(m["YIELD"] / 100.0, 6), f"{ISSB} › marketdata.YIELD (у флоатеров — при неизменном купоне)")
            if m.get("YIELD") else None,
            "as_of": "2026-09-28",
            "status": b["status"],
        })
    rep = rd["reported_2026_06_30"]
    der = rd["derived"]
    wall = [{"period": k, "bonds": V(v, calc="номинал облигаций к оферте/погашению по кварталам "
                                             "(реестр выпусков, 28.09.2026)"),
             "banks": V(None, calc="график банковских кредитов по кварталам не раскрыт")}
            for k, v in der["bonds_repayment_schedule_by_quarter_2026_09_28"].items()]
    out["debt_register"] = {
        "as_of": AS_OF,
        "collected_on": "2026-09-28",
        "unit": "млрд ₽; ставки — доли единицы",
        "basis": "финансовый долг до МСФО 16, аренда исключена",
        "anchor": {
            "borrowings_total": V(rep["borrowings_total"], f"{IFRS26} прим. 15, с. 20"),
            "bonds_carrying": V(rep["bonds_carrying"], f"{IFRS26} прим. 15, 25, с. 20, 25"),
            "bonds_face": V(der["bonds_face_2026_06_30"], calc="сумма номиналов выпусков в обращении на 30.06.2026"),
            "bank_loans": {
                "short": V(41.913, f"{IFRS26} прим. 15, с. 20 (двусторонние кредиты, погашение 2026–2027)"),
                "long": V(173.501, f"{IFRS26} прим. 15, с. 20 (двусторонние кредиты и займы, погашение 2027–2028)"),
                "total": V(215.414, calc="41,913 + 173,501"),
                "lenders": "не раскрыты", "rates": "по траншам не раскрыты",
            },
            "total_debt_company": V(435.857, f"{PR226}, с. 12; {DB26} › Debt!R6"),
            "cash": V(125.21, f"{IFRS26} прим. 8, с. 15"),
            "net_debt": V(310.647, f"{PR226}, с. 12; {DB26} › Debt!R11"),
            "floating_share": V(0.55, f"{IFRS26} прим. 24, с. 24 («55% заёмных средств с плавающей процентной "
                                      "ставкой, привязанной к ключевой ставке»)"),
            "floating_share_2025_12_31": V(0.39, f"{IFRS25} прим. 30, с. 50"),
            "effective_rate_2026H1": V(0.1654, f"{IFRS26} прим. 15, с. 20 (16,54 % годовых)"),
            "effective_rate_2025": V(0.1992, f"{IFRS25} прим. 21, с. 42"),
            "sensitivity_pbt_per_100bp_half": V(0.641, f"{IFRS26} прим. 24, с. 24", "−0,641 млрд прибыли до налога "
                                                                                    "за полугодие при +100 б.п."),
            "capitalised_interest_2026H1": V(2.085, f"{IFRS26} прим. 15, с. 20"),
            "credit_lines_unused": V(804.643, f"{IFRS26} прим. 24, с. 25"),
            "bond_programme_remaining": V(103.0, f"{IFRS26} прим. 24, с. 25"),
            "undiscounted_loans_with_interest": {
                "lt_1y": V(185.24, f"{IFRS26} прим. 24, с. 24"),
                "1_5y": V(332.86, f"{IFRS26} прим. 24, с. 24"),
                "gt_5y": V(0.0, f"{IFRS26} прим. 24, с. 24"),
            },
            "covenant": "чистый долг / EBITDA до МСФО 16 ≤ 4,00х (4,25х два квартала после приобретения); факт 1,08х "
                        f"({IFRS26} прим. 16, с. 21)",
            "weighted_avg_maturity_2025_12_31_months": V(17, "x5-ar25.pdf, с. 225"),
        },
        "ratings": [
            {"agency": "АКРА", "rating": "AAA(RU)", "outlook": "стабильный", "date": "2026-06-24",
             "src": "x5-news_akra-podtverdilo-kreditnyj-rejting-x5-na-urovne-aaaru-prognoz-stabilnyj.html"},
            {"agency": "Эксперт РА", "rating": "ruAAA", "outlook": "стабильный", "date": "2026-07-29",
             "src": "x5-news_ekspert-ra-podtverdil-kreditnyj-rejting-x5-na-urovne-ruaaa-prognoz-stabilnyj.html"},
            {"agency": "НКР", "rating": "AAA.ru", "outlook": "стабильный", "date": "2025-12-03",
             "src": "x5-ar25.pdf, с. 227 (присвоен 03.12.2025); x5-credit-ratings-2026-09-28.html"},
        ],
        "bonds": bonds,
        "bonds_summary_2026_09_28": {
            "face": V(der["bonds_face_2026_09_28"], calc="219,708 + 22,0 (003P-20) + 17,5 (003P-21) − 10,0 (003P-04)"),
            "floating_face": V(der["bonds_floating_face_2026_09_28"], calc="сумма номиналов флоатеров"),
            "floating_share": V(der["bonds_floating_share_2026_09_28_pct"] / 100.0, calc="floating_face / face"),
            "weighted_coupon_now": V(der["bonds_weighted_coupon_now_pct"] / 100.0,
                                     calc="средний купон, взвешенный по номиналу (флоатеры — КС 14,00 % + спред)"),
            "weighted_fixed_coupon": V(der["bonds_weighted_fixed_coupon_pct"] / 100.0, calc="по фиксированным выпускам"),
            "weighted_floating_spread": V(der["bonds_weighted_floating_spread_pp"] / 100.0, calc="по флоатерам"),
            "weighted_years_to_put": V(der["bonds_weighted_years_to_put"], calc="лет до оферты/погашения, взвешенно"),
            "weighted_market_yield": V(der["bonds_weighted_market_yield_pct"] / 100.0,
                                       calc="доходность ISS, взвешенная по номиналу"),
        },
        "key_rate_2026_09_28": V(0.14, "cbr-keyrate-2025-01-01_2026-09-28.html (14,00 % с 27.07.2026)"),
        "bank_floating_est": V(der["bank_floating_est_2026_06_30"],
                               calc="оценка: 55 % × 434,709 − плавающие облигации по номиналу 95,0"),
        "wall": wall,
        "wall_note": "облигации — по оферте, если её нет — по погашению (номинал на 28.09.2026); банковские кредиты "
                     "по срокам не раскрыты: краткосрочные 41,9 (окончательное погашение 2026–2027), долгосрочные "
                     "173,5 (2027–2028)",
    }

    # ================================================================ history.json
    history = {"annual": [], "halves": [], "formats": [], "format_area": []}
    # --- годы
    years = list(range(2011, 2026))
    ann, ann_cf = {}, {}
    # обесценение внеоборотных активов нетто (МСФО) — там, где собранная первичка его раскрывает
    imp_year = {2023: imp["FY2023"], 2024: imp["FY2024"], 2025: imp["FY2025"]}
    imp_hv = {"2024H1": (imp["2024H1"][0], imp["2024H1"][1], None),
              **{p: imp_half[p] for p in ("2024H2", "2025H1", "2025H2", "2026H1")}}
    DA_ALL = "D&A и обесценение внеоборотных активов до МСФО 16 (databook, SG&A стр. 8) / выручка — с обесценением"
    DA_EXCL = "(D&A и обесценение − обесценение внеоборотных активов нетто по МСФО) / выручка — без обесценения"
    NO_IMP = "обесценение за период в собранной первичке не раскрыто (МСФО есть с 2023 г.)"
    OIP = ("«Прочие платежи по инвестиционной деятельности» (ОДДС), отток — положительным; в capex (ОС + НМА), "
           "M&A и мост не входят")
    FLR = "«Поступления от основной суммы чистых инвестиций в аренду» (ОДДС: финансовая аренда франчайзи)"

    def cf_node(x, s, sign, what):
        """Строка ОДДС в млрд ₽ (пустая ячейка databook — строки в периоде нет)."""
        if x is None:
            return V(None, calc=f"строки в ОДДС за период нет — появилась в 4 кв. 2023 г. ({s})")
        return V(bn(sign * x) + 0.0, s, what)

    for y in years:
        if y <= 2020:
            bk, cm = b24, c24
            colm = {s: (cm[s].get(str(y)) or cm[s].get(y)) for s in cm}
            rev = bk.val("Profit and Loss", f"{colm['Profit and Loss']}6")
            adj = bk.val("EBITDA", f"{colm['EBITDA']}14")
            ebr = bk.val("EBITDA", f"{colm['EBITDA']}18")
            da = bk.val("SG&A", f"{colm['SG&A']}8")
            cx = -(bk.val("Cash Flow", f"{colm['Cash Flow']}33") + bk.val("Cash Flow", f"{colm['Cash Flow']}42"))
            lev = bk.val("Debt", f"{colm['Debt']}12")
            r_oip, r_flr = 47, 46
            refs = {"rev": bk.ref("Profit and Loss", f"{colm['Profit and Loss']}6"),
                    "adj": bk.ref("EBITDA", f"{colm['EBITDA']}14"), "ebr": bk.ref("EBITDA", f"{colm['EBITDA']}18"),
                    "da": bk.ref("SG&A", f"{colm['SG&A']}8"),
                    "cx": bk.ref("Cash Flow", f"{colm['Cash Flow']}33", f"{colm['Cash Flow']}42"),
                    "lev": bk.ref("Debt", f"{colm['Debt']}12")}
        else:
            bk, cm = b26, c26
            colm = {s: (cm[s].get(str(y)) or cm[s].get(y)) for s in cm}
            rev = bk.val("Profit and Loss", f"{colm['Profit and Loss']}6")
            adj = bk.val("EBITDA", f"{colm['EBITDA']}14")
            ebr = bk.val("EBITDA", f"{colm['EBITDA']}22")
            da = bk.val("SG&A", f"{colm['SG&A']}8")
            cx = -(bk.val("Cash Flow", f"{colm['Cash Flow']}30") + bk.val("Cash Flow", f"{colm['Cash Flow']}36"))
            dcol = cm["Debt"][f"{y}-12-31"]
            lev = bk.val("Debt", f"{dcol}12")
            r_oip, r_flr = 42, 41
            refs = {"rev": bk.ref("Profit and Loss", f"{colm['Profit and Loss']}6"),
                    "adj": bk.ref("EBITDA", f"{colm['EBITDA']}14"), "ebr": bk.ref("EBITDA", f"{colm['EBITDA']}22"),
                    "da": bk.ref("SG&A", f"{colm['SG&A']}8"),
                    "cx": bk.ref("Cash Flow", f"{colm['Cash Flow']}30", f"{colm['Cash Flow']}36"),
                    "lev": bk.ref("Debt", f"{dcol}12")}
        # прочие инвестиционные платежи и поступления по финансовой аренде (ОДДС; в capex, M&A и мост не входят)
        cfc = colm["Cash Flow"]
        refs["oip"], refs["flr"] = bk.ref("Cash Flow", f"{cfc}{r_oip}"), bk.ref("Cash Flow", f"{cfc}{r_flr}")
        ann_cf[y] = (bk.val("Cash Flow", f"{cfc}{r_oip}"), bk.val("Cash Flow", f"{cfc}{r_flr}"))
        ann[y] = (rev, adj, ebr, da, cx, lev, refs)
    # LFL, площадь, магазины по годам
    orq = c24["Operating Results_Q"]

    def oldop(row, y):
        col = orq.get(str(y)) or orq.get(y)
        return b24.val("Operating Results_Q", f"{col}{row}"), b24.ref("Operating Results_Q", f"{col}{row}")

    for y in years:
        rev, adj, ebr, da, cx, lev, refs = ann[y]
        prev = ann.get(y - 1)
        row = {"year": y,
               "revenue": V(bn(rev), refs["rev"]),
               "growth": V(r6(rev / prev[0] - 1), calc=f"выручка {y} / {y - 1} − 1"
                           + (" (2021 г. — новый databook к старому)" if y == 2021 else "")) if prev else V(None, calc="нет 2010 г."),
               "adj_margin": V(r6(adj / rev), refs["adj"] + " / " + refs["rev"].split(" › ")[1], "скорр. EBITDA / выручка"),
               "rep_margin": V(r6(ebr / rev), refs["ebr"] + " / " + refs["rev"].split(" › ")[1], "EBITDA / выручка"),
               "capex_pct": V(r6(cx / rev), refs["cx"], "(ОС + НМА) / выручка"),
               "da_pct": V(r6(da / rev), refs["da"], DA_ALL),
               "da_excl_impairment_pct": (V(r6((da - imp_year[y][0]) / rev), refs["da"] + "; " + imp_year[y][1],
                                            DA_EXCL + f": обесценение {ru(bn(imp_year[y][0]))}")
                                          if y in imp_year else V(None, calc=NO_IMP)),
               "other_investing_payments": cf_node(ann_cf[y][0], refs["oip"], -1.0, OIP),
               "finance_lease_receipts": cf_node(ann_cf[y][1], refs["flr"], 1.0, FLR),
               "leverage": V(r6(lev), refs["lev"], "чистый долг / EBITDA до МСФО 16 на 31.12"),
               }
        if y <= 2023:
            for key, rr in (("lfl", 45), ("lfl_traffic", 46), ("lfl_ticket", 47)):
                v, s = oldop(rr, y)
                row[key] = V(r6(v), s)
        else:
            col = orc["2024"] if y == 2024 else orc["2025 (2)"]
            for key, rr in (("lfl", 28), ("lfl_traffic", 29), ("lfl_ticket", 30)):
                row[key] = V(r6(b26.val(OR, f"{col}{rr}")), b26.ref(OR, f"{col}{rr}"))
        if y <= 2022:
            va, sa = oldop(76, y)
            vs, ss = oldop(68, y)
            row["area_end"] = V(r6(va), sa, "итого X5 как отчитано (с 2022 г. — с «Красным Яром»/«Слатой»)")
            row["stores_end"] = V(vs, ss, "итого X5 как отчитано")
        else:
            p = f"{y}H2"
            row["area_end"] = V(area_end[p]["v"], calc=f"network.area_end[{p}] — сопоставимый базис")
            row["stores_end"] = V(stores_end_all[p]["v"], calc=f"network.stores_end_hist[{p}] — сопоставимый базис")
        history["annual"].append(row)

    # форматы (чистая розничная выручка с экспресс-доставкой), млрд ₽
    for y in years:
        rev = ann[y][0]
        if y <= 2023:
            vals = {}
            for key, rr in (("pyaterochka", 7), ("perekrestok", 9), ("karusel", 11), ("chizhik", 13)):
                v, s = oldop(rr, y)
                vals[key] = (v, s)
            dg = oldop(39, y) if y >= 2019 else (None, None)
            area_rows = (("pyaterochka", 70), ("perekrestok", 71), ("karusel", 72), ("chizhik", 73))
            areas = {k: oldop(rr, y) for k, rr in area_rows}
            if y == 2023:   # площадь 31.12.2023 — в сопоставимом базисе (КЯ/Слата в «Пятёрочке»)
                areas.update({f: (area_fmt["2023H2"][f][0], area_fmt["2023H2"][f][1]) for f in fmts})
        else:
            col = orc["2024"] if y == 2024 else orc["2025 (2)"]   # 2024 — столбец «Скорр.*» (E)
            vals = {"pyaterochka": (b26.val(OR, f"{col}8"), b26.ref(OR, f"{col}8")),
                    "perekrestok": (b26.val(OR, f"{col}10"), b26.ref(OR, f"{col}10")),
                    "chizhik": (b26.val(OR, f"{col}12"), b26.ref(OR, f"{col}12")),
                    "karusel": (0.0, "формат закрыт в 2023 г.")}
            dg = (b26.val(OR, f"{col}20"), b26.ref(OR, f"{col}20"))
            areas = {"pyaterochka": (area_fmt[f"{y}H2"]["pyaterochka"][0], area_fmt[f"{y}H2"]["pyaterochka"][1]),
                     "perekrestok": (area_fmt[f"{y}H2"]["perekrestok"][0], area_fmt[f"{y}H2"]["perekrestok"][1]),
                     "chizhik": (area_fmt[f"{y}H2"]["chizhik"][0], area_fmt[f"{y}H2"]["chizhik"][1]),
                     "karusel": (0.0, "формат закрыт в 2023 г.")}
        frow = {"year": y}
        tot = 0.0
        for k in ("pyaterochka", "perekrestok", "chizhik", "karusel"):
            v, s = vals[k]
            frow[k] = V(bn(v) if v is not None else None, s) if v is not None else V(None, calc="формата нет")
            tot += v or 0.0
        frow["other"] = V(bn(rev - tot), calc="выручка − сумма форматов (цифровые без экспресс-доставки, опт, "
                                              "франшиза, прочие бизнесы; до 2024 г. — и «Красный Яр»/«Слата»)")
        frow["digital"] = V(bn(dg[0]), dg[1], "справочно: выручка цифровых бизнесов (пересекается с форматами — "
                                              "экспресс-доставка)") if dg[0] is not None else V(None, calc="не раскрыто")
        history["formats"].append(frow)
        arow = {"year": y}
        for k in ("pyaterochka", "perekrestok", "chizhik", "karusel"):
            v, s = areas[k]
            arow[k] = V(r6(v), s) if v is not None else V(None, calc="формата нет")
        history["format_area"].append(arow)

    # --- полугодия 2018H1–2026H1 (+ 2017 для роста 2018)
    halves = [f"{y}H{h}" for y in range(2017, 2027) for h in (1, 2)][:-1]
    hv = {}
    for p in halves:
        rev, s_rev = half("Profit and Loss", 6, "Profit and Loss", 6, p)
        if p.startswith("2017"):
            hv[p] = (rev, s_rev)
            continue
        adj, s_adj = half("EBITDA", 14, "EBITDA", 14, p)
        ebr, s_ebr = half("EBITDA", 22, "EBITDA", 18, p)
        da, s_da = half("SG&A", 8, "SG&A", 8, p)
        cxa, s_cxa = half("Cash Flow", 30, "Cash Flow", 33, p, sign=-1.0)
        cxb, s_cxb = half("Cash Flow", 36, "Cash Flow", 42, p, sign=-1.0)
        oip, s_oip = half("Cash Flow", 42, "Cash Flow", 47, p)
        flr, s_flr = half("Cash Flow", 41, "Cash Flow", 46, p)
        hv[p] = (rev, s_rev)
        prev = f"{int(p[:4]) - 1}{p[4:]}"
        pr = hv[prev][0]
        row = {"period": p,
               "revenue": V(bn(rev), s_rev, "сумма кварталов"),
               "growth": V(r6(rev / pr - 1), calc=f"выручка {p} / {prev} − 1"),
               "adj_margin": V(r6(adj / rev), s_adj, "скорр. EBITDA / выручка"),
               "rep_margin": V(r6(ebr / rev), s_ebr, "EBITDA / выручка"),
               "capex_pct": V(r6((cxa + cxb) / rev), s_cxa + "; " + s_cxb, "(ОС + НМА) / выручка"),
               "da_pct": V(r6(da / rev), s_da, DA_ALL),
               "da_excl_impairment_pct": (V(r6((da - imp_hv[p][0]) / rev), s_da + "; " + imp_hv[p][1],
                                            DA_EXCL + f": обесценение {ru(bn(imp_hv[p][0]))}"
                                            + (f" ({imp_hv[p][2]})" if imp_hv[p][2] else ""))
                                          if p in imp_hv else V(None, calc=NO_IMP)),
               "other_investing_payments": cf_node(oip, s_oip, -1.0, OIP),
               "finance_lease_receipts": cf_node(flr, s_flr, 1.0, FLR),
               }
        y, h = int(p[:4]), int(p[-1])
        if y >= 2023 and p >= "2023H2":
            if p in area_end:
                row["area_end"] = V(area_end[p]["v"], calc=f"network.area_end[{p}] — сопоставимый базис")
        else:
            key = q_old(2 if h == 1 else 4, y)
            col = orq[key]
            row["area_end"] = V(r6(b24.val("Operating Results_Q", f"{col}76")),
                                b24.ref("Operating Results_Q", f"{col}76"), "итого X5 как отчитано")
        # рычаг на конец полугодия
        if p >= "2023H2":
            dkey = {"2023H2": "2023-12-31", "2024H1": "2024-06-30", "2024H2": "2024-12-31", "2025H1": "2025-06-30",
                    "2025H2": "2025-12-31", "2026H1": "2026-06-30"}[p]
            dcol = c26["Debt"][dkey]
            row["leverage"] = V(r6(b26.val("Debt", f"{dcol}12")), b26.ref("Debt", f"{dcol}12"))
        else:
            dcol = c24["Debt"][q_old(2 if h == 1 else 4, y)]
            lv = b24.val("Debt", f"{dcol}12")
            row["leverage"] = V(r6(lv), b24.ref("Debt", f"{dcol}12")) if lv is not None else V(None, calc="не раскрыто")
        history["halves"].append(row)
    out["history"] = {
        "unit": "выручка — млрд ₽; доли — доли единицы; площадь — тыс. м²",
        "basis": "до МСФО 16; годы ≤2020 и полугодия ≤2022H2 — databook X5 Retail Group N.V. "
                 f"({DB24}), дальше — databook ПАО «КЦ ИКС 5» ({DB26}); площадь и магазины с 2023H2 — "
                 "сопоставимый базис network.json",
        **history,
    }

    # ================================================================ peers.json
    fpb = Book(PRIMARY / FP, FP)
    fp_ebitda = fpb.val("5", "G23") - fpb.val("5", "M23") + fpb.val("5", "N23")
    fp_np = fpb.val("1", "G22") - fpb.val("1", "M22") + fpb.val("1", "N22")
    fp_rev = fpb.val("1", "G6") - fpb.val("1", "M6") + fpb.val("1", "N6")
    lb = Book(LENTA_DB, LENTA_DB.name).wb
    lpl, lbs = lb["PL"], lb["BS"]
    l_nd = (lbs["AF54"].value + lbs["AF62"].value - lbs["AF34"].value) / 1e6
    l_eb = (lpl["AD49"].value - lpl["AC49"].value + lpl["AE49"].value) / 1e6
    l_np = (lpl["AT46"].value - lpl["AS46"].value + lpl["AU46"].value) / 1e6
    l_np17 = (lpl["AD46"].value - lpl["AC46"].value + lpl["AE46"].value) / 1e6
    l_rev = (lpl["AD10"].value - lpl["AC10"].value + lpl["AE10"].value) / 1e6
    LEN = f"lenta-850-handoff/reference/primary/Lenta_Q22026_DATABOOK.xlsx (sha256 {sha256_file(LENTA_DB)[:16]}…)"
    mh = json.loads((MAGNIT / "data/facts/history.json").read_text(encoding="utf-8"))
    mpre, mifrs = mh["pnl"]["pre_ifrs16"], mh["pnl"]["ifrs16"]
    m_eb = (mpre["2025H2"]["ebitda"] + mpre["2026H1"]["ebitda"]) / 1e6
    m_rev = (mpre["2025H2"]["revenue"] + mpre["2026H1"]["revenue"]) / 1e6
    m_np = (mifrs["2025H2"]["net_income"] + mifrs["2026H1"]["net_income"]) / 1e6
    m_np17 = (mpre["2025H2"]["net_income"] + mpre["2026H1"]["net_income"]) / 1e6
    mnd = json.loads((MAGNIT / "data/assumptions/evidence/book-1.4/beta/inputs/net_debt.json").read_text(encoding="utf-8"))
    MAG = "magnit-850oa (ветка v5, только чтение)"
    # чистая прибыль X5 LTM по МСФО 16 (P&L строка 22, столбцы AM, AN, AO, AP)
    x5_np_cells = ["AM22", "AN22", "AO22", "AP22"]
    x5_np = sum(b26.val("Profit and Loss", c) for c in x5_np_cells) / 1000.0
    x5_np17_cells = ["Y22", "Z22", "AA22", "AB22"]
    x5_np17 = sum(b26.val("Profit and Loss", c) for c in x5_np17_cells) / 1000.0
    out["peers"] = {
        "as_of": AS_OF,
        "unit": "млрд ₽; акции — млн шт.",
        "basis": "одна база для всех: чистый долг до МСФО 16 на 30.06.2026, EBITDA до МСФО 16 (отчётная) и чистая "
                 "прибыль за 12 мес. на 30.06.2026 (LTM 1П2026); цена — живая, берёт конвейер",
        "rows": [
            {"ticker": "X5", "name": "X5",
             "shares_outstanding": V(245.98118, f"{IFRS26} прим. 17, с. 21"),
             "net_debt": V(310.647, f"{DB26} › Debt!R11"),
             "dividends_after_balance": V(60.265, f"{IFRS26} прим. 17, с. 22",
                                          "объявлены до даты баланса, реестр 07.07.2026 — после"),
             "ebitda_ltm": V(r6(rep_ltm), calc="balance.ebitda_rep_ltm"),
             "revenue_ltm": V(r6(rev_ltm), calc="balance.revenue_ltm"),
             "net_profit_ltm": V(round(x5_np, 6), b26.ref("Profit and Loss", *x5_np_cells),
                                 "чистая прибыль по МСФО (IFRS 16), 3 кв. 2025 – 2 кв. 2026"),
             "net_profit_ltm_pre16": V(round(x5_np17, 6), b26.ref("Profit and Loss", *x5_np17_cells)),
             "reported_on": "2026-08-13"},
            {"ticker": "MGNT", "name": "Магнит",
             "shares_outstanding": V(67.847, f"{MAG} data/assumptions/assumptions.yaml › facts.shares_mln (в обращении; "
                                             "выпущено 101,911, у дочерних 34,064)"),
             "net_debt": V(round(mnd["MGNT"]["2026-06-30"]["nd_bn"], 6),
                           f"{MAG} data/assumptions/evidence/book-1.4/beta/inputs/net_debt.json › MGNT.2026-06-30 "
                           f"({mnd['MGNT']['2026-06-30']['src']})"),
             "dividends_after_balance": V(0.0, f"{MAG} data/facts/peers.json (выплат нет)"),
             "ebitda_ltm": V(round(m_eb, 6), f"{MAG} data/facts/history.json › pnl.pre_ifrs16.2025H2.ebitda + "
                                             "2026H1.ebitda (databook Магнита, IAS 17)"),
             "revenue_ltm": V(round(m_rev, 6), f"{MAG} data/facts/history.json › pnl.pre_ifrs16 2025H2 + 2026H1"),
             "net_profit_ltm": V(round(m_np, 6), f"{MAG} data/facts/history.json › pnl.ifrs16.2025H2.net_income + "
                                                 "2026H1.net_income (МСФО 16)"),
             "net_profit_ltm_pre16": V(round(m_np17, 6), f"{MAG} data/facts/history.json › pnl.pre_ifrs16 "
                                                         "2025H2 + 2026H1 net_income"),
             "reported_on": "2026-08-28",
             "reported_on_src": f"{MAG} data/facts/sources.json › RELEASE_H1_2026.publication_date"},
            {"ticker": "LENT", "name": "Лента",
             "shares_outstanding": V(115.985197, f"{MAG} data/assumptions/evidence/book-1.4/beta/inputs/shares.json "
                                                 "(ISS LENT ISSUESIZE 115 985 197, запрос 24.09.2026)"),
             "net_debt": V(round(l_nd, 6), f"{LEN} › BS!AF54 + AF62 − AF34 (1П2026, IAS 17: кредиты долгосрочные + "
                                           "краткосрочные − денежные средства)"),
             "dividends_after_balance": V(0.0, f"{MAG} data/facts/peers.json («дивидендов нет»)"),
             "ebitda_ltm": V(round(l_eb, 6), f"{LEN} › PL!AD49 − AC49 + AE49 (IAS 17)"),
             "revenue_ltm": V(round(l_rev, 6), f"{LEN} › PL!AD10 − AC10 + AE10"),
             "net_profit_ltm": V(round(l_np, 6), f"{LEN} › PL!AT46 − AS46 + AU46 (МСФО 16)"),
             "net_profit_ltm_pre16": V(round(l_np17, 6), f"{LEN} › PL!AD46 − AC46 + AE46 (IAS 17)"),
             "reported_on": "2026-08-03",
             "reported_on_src": "lenta-850-handoff/reference/primary/MANIFEST.md (пресс-релиз 2 кв. 2026 от 03.08.2026)",
             "note": "О'КЕЙ консолидирован с 02.06.2026: его долг в ЧД целиком, EBITDA — за один месяц; "
                     "мультипликатор LTM завышен (research/X4 §3.1)"},
            {"ticker": "FIXR", "name": "Fix Price (ПАО)",
             "shares_outstanding": V(99856.1, "peer-fixprice-pjsc-q2-2026-results-rus.pdf, с. 5 (выкуплено 300 млн, "
                                             "передано по LTIP 250 млн, второй этап — 93,9 млн на 26.08.2026); "
                                             "ISS FIXR ISSUESIZE 100 000 000 000 (research/X4)",
                                     "100 000 − (300 − 250 + 93,9) млн акций; разница с выпущенными 0,14 %"),
             "net_debt": V(round(fpb.val("6", "N17") / 1000.0, 6), f"{FP} › 6!N17 («Скорр. чистый долг по МСФО "
                                                                   "(IAS) 17», 30.06.2026)"),
             "dividends_after_balance": V(0.0, "peer-fixprice-pjsc-q2-2026-results-rus.pdf (дивиденды за 2025 г. "
                                              "не выплачиваются)"),
             "ebitda_ltm": V(round(fp_ebitda / 1000.0, 6), f"{FP} › 5!G23 − M23 + N23 («EBITDA по МСФО (IAS) 17»)"),
             "revenue_ltm": V(round(fp_rev / 1000.0, 6), f"{FP} › 1!G6 − M6 + N6"),
             "net_profit_ltm": V(round(fp_np / 1000.0, 6), f"{FP} › 1!G22 − M22 + N22 (МСФО 16)"),
             "net_profit_ltm_pre16": V(None, calc="не раскрыта"),
             "reported_on": "2026-08-27"},
            {"ticker": "OKEY", "name": "О'Кей («ДА!»)",
             "shares_outstanding": V(269.074, "peer-okey-1h2026-ifrs-fs.pdf, с. 21 (средневзвешенное число акций "
                                              "269 074 тыс.)"),
             "net_debt": V(-8.987381, "peer-okey-1h2026-ifrs-fs.pdf, с. 3–4 (кредитов и займов нет; денежные "
                                      "средства 8 987 381 тыс. ₽)", "0 − денежные средства (чистая касса)"),
             "dividends_after_balance": V(None, calc="не раскрыто"),
             "ebitda_ltm": V(None, calc="EBITDA до МСФО 16 не раскрыта (только МСФО 16: 6М2026 4,152 млрд)"),
             "revenue_ltm": V(None, calc="LTM на одной базе не собрана: гипермаркеты проданы в ноябре 2025 г."),
             "net_profit_ltm": V(None, calc="за 2025 г. на базе продолжающейся деятельности не найдена"),
             "net_profit_ltm_pre16": V(None, calc="не раскрыта"),
             "reported_on": "2026-08-21",
             "note": "после продажи гипермаркетов осталась сеть дискаунтеров «ДА!»; торгуются ГДР (OKEY)"},
        ],
    }

    # ================================================================ brokers.json
    brokers = [
        ("2026-09-24", "Финам", "покупать", 2510, "12 мес.", "6713775"),
        ("2026-09-11", "Эйлер", "покупать", 2600, "12 мес.", "6704753"),
        ("2026-08-14", "БКС Мир инвестиций", "позитивно", 2500, "12 мес.", "6685670"),
        ("2026-08-14", "ПСБ", "потенциал роста", 3800, "долгосрочный", "6685730"),
        ("2026-08-13", "Т-Инвестиции", "держать", 2250, "12 мес.", "6685040"),
        ("2026-08-13", "Цифра брокер", "покупать", 3943, None, "6685096"),
        ("2026-08-13", "Синара", "позитивно", None, None, "6684897"),
        ("2026-08-13", "Газпромбанк", "позитивно на долгий срок", None, None, "6685216"),
        ("2026-05-19", "Freedom Finance Global", "держать", 3000, None, "6623908"),
        ("2026-04-30", "ВЕЛЕС Капитал", "покупать", 3603, None, "6611603"),
        ("2026-04-30", "ВТБ Мои инвестиции", "позитивно", 3485, "12 мес.", "6611561"),
    ]
    rows = []
    for d, name, rating, tgt, hor, fid in brokers:
        f = f"analyst-finmarket-{fid}.html"
        rows.append({"broker": name, "date": d, "rating": rating,
                     "target": V(tgt, f"{f} (Интерфакс/Финмаркет; research/X4 §4)") if tgt is not None
                     else V(None, calc=f"цель не раскрыта ({f})"),
                     "horizon": hor, "src": f})
    t12 = sorted(r["target"]["v"] for r in rows if r["target"]["v"] and r["horizon"] == "12 мес."
                 and r["date"] >= "2026-08-13")
    mid = (t12[len(t12) // 2 - 1] + t12[len(t12) // 2]) / 2 if len(t12) % 2 == 0 else t12[len(t12) // 2]
    out["brokers"] = {
        "as_of": "2026-09-28",
        "unit": "₽ на акцию",
        "rows": rows,
        "median_12m_after_2q2026": V(mid, calc="медиана 12-месячных целей после отчёта за 2 кв. 2026 (13.08–24.09.2026): "
                                               + " / ".join(str(x) for x in t12)),
        "after_report": "МСФО 2 кв. 2026 (13.08.2026)",
        "not_included": "вторичные и устаревшие: консенсус 3 290 ₽ (smart-lab 22.07.2026, источник не назван), "
                        "SberCIB 4 500 ₽ (пересказы, декабрь 2025), ВТБ Моя аналитика 4 481 ₽ (17.10.2025); "
                        "Альфа, Атон, Совкомбанк, Invest Heroes — актуальной цели с датой нет (research/X4 §4)",
    }

    # ================================================================ actuals.json
    out["actuals"] = {
        "description": "Факты по целям журнала после отчётов; вносятся человеком после публикации отчёта X5. "
                       "Формат записи: {\"target\": \"x5.adj_margin\" | \"x5.revenue_growth\", \"period\": "
                       "\"2026H2\", \"value\": {\"v\": 0.061, \"src\": \"пресс-релиз 4 кв. 2026, с. 3\"}, "
                       "\"reported_on\": \"2027-03-19\"}. value — узел факта, как во всех файлах фактов: "
                       "{\"v\": доля единицы, \"src\": \"файл первички › место\"} (или \"calc\" вместо "
                       "\"src\"); число без узла — число без источника, журнал его не принимает (отказ "
                       "сборки). Одна запись на цель и полугодие. adj_margin — скорр. EBITDA до МСФО 16 / "
                       "выручка полугодия; revenue_growth — рост выручки полугодия г/г.",
        "actuals": [],
    }

    # ================================================================ guidance.json
    G = "x5-news-2026-03-20-guidance-2026.html (пресс-релиз «X5 представила прогнозы и ориентиры на 2026 год», 20.03.2026)"
    out["guidance"] = {
        "year": 2026,
        "published": "2026-03-20",
        "revenue_growth": V([0.12, 0.16], G, "«рост выручки на уровне 12–16%»"),
        "adj_margin_min": V(0.06, G, "«рентабельность скорректированной EBITDA на уровне не ниже 6%» (до МСФО 16)"),
        "capex_pct": V([0.045, 0.047], G, "«капитальные затраты на уровне 4,5–4,7% от выручки»"),
        "openings_min": V(2000, f"{PR226}, с. 2 (письмо CEO: «план по открытию более 2 000 магазинов в 2026 году»)",
                          "магазинов за 2026 г. (подтверждено 13.08.2026)"),
        "net_debt_to_ebitda": V([1.2, 1.4], G, "«поддержание целевого показателя „чистый долг/EBITDA“ на уровне "
                                               "1,2-1,4х»"),
        "base_2025": {"revenue": V(4642.034, f"{DB26} › Profit and Loss!H6"),
                      "adj_ebitda": V(285.467649, f"{DB26} › EBITDA!H14")},
        "implied_2026": {
            "revenue": V([round(4642.034 * 1.12, 3), round(4642.034 * 1.16, 3)], calc="выручка 2025 × (1,12; 1,16)"),
            "revenue_2026H2_growth": V([r6((4642.034 * 1.12 - periods["2026H1"]["revenue"]["v"]) /
                                           periods["2025H2"]["revenue"]["v"] - 1),
                                        r6((4642.034 * 1.16 - periods["2026H1"]["revenue"]["v"]) /
                                           periods["2025H2"]["revenue"]["v"] - 1)],
                                       calc="(выручка 2026 по прогнозу − 1П2026) / 2П2025 − 1"),
            "adj_margin_2026H2_min": V(r6((0.06 * 4642.034 * 1.12 - periods["2026H1"]["adj_ebitda"]["v"]) /
                                          (4642.034 * 1.12 - periods["2026H1"]["revenue"]["v"])),
                                       calc="(6 % × выручка 2026 при росте 12 % − скорр. EBITDA 1П2026) / выручка 2П2026"),
        },
    }

    # ================================================================ calendar.json
    # У МСФО седьмой элемент — `covers`: полугодие, которое отчёт закрывает (годовое — H2,
    # за 2 кв. и 1П — H1), None — МСФО за 1/3 кв. внутри полугодия (data/facts/SCHEMA.md,
    # «calendar.json»). У остальных событий поля нет.
    CAL = "closed-periods-calendar-2026_rus.pdf, с. 3"
    CBR = "cbr-cal-mp-2026-09-28.html (cbr.ru/dkp/cal_mp/, доступ 28.09.2026)"
    PREC = "research/indicators/x5_calendar.json"
    ev = [
        ("2026-10-16", "Операционные результаты X5 за 3 кв. 2026 г.", "trading_update", True, "", CAL),
        ("2026-10-23", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", True,
         "среднесрочный прогноз", CBR),
        ("2026-10-29", "Финансовые результаты X5 за 3 кв. 2026 г. (МСФО)", "ifrs", True, "", CAL, None),
        ("2026-11-13", "Рекомендация Наблюдательного совета по дивидендам за 9 мес. 2026 г.", "dividend", False,
         "оценка по прецеденту: 13.11.2025 (≈ середина ноября)", PREC),
        ("2026-12-18", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", True, "", CBR),
        ("2026-12-18", "ВОСА о дивидендах за 9 мес. 2026 г.", "dividend", False,
         "оценка по прецеденту: ВОСА 18.12.2025", PREC),
        ("2027-01-06", "Дата закрытия реестра (отсечка) по дивидендам за 9 мес. 2026 г.", "dividend", False,
         "оценка по прецеденту: реестр 06.01.2026 (≈ начало января 2027)", PREC),
        ("2027-01-28", "Операционные результаты X5 за 4 кв. и 2026 г.", "trading_update", False,
         "оценка по прецеденту: 24.01.2024, 27.01.2025, 28.01.2026 (≈ конец января)", PREC),
        ("2027-02-12", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: календарь ЦБ на 2027 г. не опубликован на 28.09.2026; прецедент 16.02.2024, 14.02.2025, 13.02.2026",
         CBR),
        ("2027-03-19", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: прецедент 22.03.2024, 21.03.2025, 20.03.2026", CBR),
        ("2027-03-19", "Финансовые результаты X5 за 2026 г. (МСФО) и ориентиры на 2027 г.", "ifrs", False,
         "оценка по прецеденту: пятница около 20 марта — 22.03.2024, 21.03.2025, 20.03.2026 (≈ 19.03.2027)", PREC,
         "2026H2"),
        ("2027-04-16", "Операционные результаты X5 за 1 кв. 2027 г.", "trading_update", False,
         "оценка по прецеденту: 16.04.2024, 16.04.2025, 16.04.2026", PREC),
        ("2027-04-23", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: прецедент 26.04.2024, 25.04.2025, 24.04.2026", CBR),
        ("2027-04-29", "Финансовые результаты X5 за 1 кв. 2027 г. (МСФО)", "ifrs", False,
         "оценка по прецеденту: 22.04.2024, 05.05.2025, 29.04.2026", PREC, None),
        ("2027-05-19", "Рекомендация Наблюдательного совета по дивидендам за 2026 г.", "dividend", False,
         "оценка по прецеденту: 19.05.2026", PREC),
        ("2027-06-11", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: прецедент 07.06.2024, 06.06.2025, 19.06.2026", CBR),
        ("2027-06-25", "Годовое общее собрание акционеров X5 (дивиденды за 2026 г.)", "dividend", False,
         "оценка по прецеденту: ГОСА 27.06.2025, 26.06.2026 (≈ конец июня)", PREC),
        ("2027-07-07", "Дата закрытия реестра по финальным дивидендам за 2026 г.", "dividend", False,
         "оценка по прецеденту: 09.07.2025, 07.07.2026", PREC),
        ("2027-07-16", "Операционные результаты X5 за 2 кв. 2027 г.", "trading_update", False,
         "оценка по прецеденту: 16.07.2024, 16.07.2025, 16.07.2026", PREC),
        ("2027-07-23", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: прецедент 26.07.2024, 25.07.2025, 24.07.2026", CBR),
        ("2027-08-13", "Финансовые результаты X5 за 2 кв. и 1П 2027 г. (МСФО)", "ifrs", False,
         "оценка по прецеденту: 15.08.2024, 13.08.2025, 13.08.2026", PREC, "2027H1"),
        ("2027-09-10", "Заседание Совета директоров Банка России по ключевой ставке", "cbr", False,
         "оценка: прецедент 13.09.2024, 12.09.2025, 11.09.2026", CBR),
    ]
    out["_calendar"] = {
        "as_of": "2026-09-28",
        "horizon": "ближайшие 12 месяцев",
        "events": [{"date": d, "title": t, "kind": k, **({"covers": cov[0]} if k == "ifrs" else {}),
                    "confirmed": c, "note": n, "src": s}
                   for d, t, k, c, n, s, *cov in sorted(ev, key=lambda e: (e[0], e[2]))],
    }
    return out


def dump(data) -> bytes:
    return (json.dumps(data, ensure_ascii=False, indent=1) + "\n").encode("utf-8")


def write(out: dict, root: Path) -> list:
    """Пишет root/data/facts/*.json и root/data/calendar.json; возвращает записанные пути."""
    out = dict(out)
    out.pop("_stores_check", None)
    cal = out.pop("_calendar")
    # actuals.json ведёт человек после отчётов (docs/INDICATORS.md): сборщик задаёт только
    # формат и не перезаписывает внесённые факты. Пишем его, лишь если файла ещё нет.
    actuals = out.pop("actuals")
    facts = root / "data" / "facts"
    facts.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, data in out.items():
        paths.append(facts / f"{name}.json")
        paths[-1].write_bytes(dump(data))
    paths.append(root / "data" / "calendar.json")
    paths[-1].write_bytes(dump(cal))
    manual = facts / "actuals.json"
    if not manual.exists():
        manual.write_bytes(dump(actuals))
    return paths


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--primary", help="папка первички (по умолчанию X5_PRIMARY или "
                                      "../x5-850-handoff/reference/primary)")
    ap.add_argument("--workspace", help="рабочий каталог с x5-850-handoff, magnit-850oa, lenta-850-handoff "
                                        "(по умолчанию X5_WORKSPACE или папка над репозиторием)")
    ap.add_argument("--out", default=str(REPO), help="корень, куда писать data/ (по умолчанию репозиторий)")
    args = ap.parse_args(argv)
    src = sources(args.workspace, args.primary)
    lost = missing(src)
    if lost:
        raise SystemExit("нет входов сборки: " + "; ".join(lost))
    out = build(src)
    bad = {k: v for k, v in out["_stores_check"].items() if v[0] != v[1]}
    paths = write(out, Path(args.out))
    print("записано:", ", ".join(p.name for p in paths), "→", Path(args.out) / "data")
    print("сверка чистых открытий магазинов (кварталы vs концы полугодий), расхождения:", bad or "нет")


if __name__ == "__main__":
    main()
